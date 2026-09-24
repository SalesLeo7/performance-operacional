"""Prepara os dados iniciais usados para criar a base de parâmetros operacionais."""

from __future__ import annotations

import json
import math
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "00 - Bases OVERVIEW HTMLs.xlsx"
SLA_CONFIG = ROOT / "MetasSLA.xlsx"
OUTPUT = Path(__file__).resolve().parent / "data" / "operational_defaults.json"


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    raw = clean(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            pass
    return None


def parse_time(value: Any) -> time | None:
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    if isinstance(value, (int, float)):
        seconds = round(float(value) * 86400) % 86400
        return time(seconds // 3600, seconds % 3600 // 60, seconds % 60)
    raw = clean(value)
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(raw, fmt).time()
        except ValueError:
            pass
    return None


def parse_lines(value: Any) -> int:
    try:
        return max(0, int(round(float(value))))
    except (TypeError, ValueError):
        return 0


def percentile(values: list[int], quantile: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    position = (len(ordered) - 1) * quantile
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return float(ordered[lower])
    return ordered[lower] + (ordered[upper] - ordered[lower]) * (position - lower)


def next_weekday(day: date) -> date:
    while day.weekday() >= 5:
        day += timedelta(days=1)
    return day


def release_business_day(stamp: datetime) -> date:
    day = next_weekday(stamp.date())
    if day != stamp.date() or stamp.time() >= time(17):
        return next_weekday(day + timedelta(days=1))
    return day


def aliases() -> tuple[dict[str, dict[str, str]], list[dict[str, str]]]:
    mapping: dict[str, dict[str, str]] = {}
    clients: dict[str, str] = {}
    book = load_workbook(SLA_CONFIG, read_only=True, data_only=True)
    rows = book["ClientAliases"].iter_rows(values_only=True)
    headers = {clean(value): index for index, value in enumerate(next(rows))}
    for row in rows:
        alias = clean(row[headers["ClientAlias"]]).upper()
        key = clean(row[headers["ClientKey"]]).upper()
        name = clean(row[headers["ClientName"]]) or key
        operation = clean(row[headers["Operation"]])
        if alias:
            mapping[alias] = {"key": key, "name": name, "operation": operation}
        if operation in {"Outbound", "Both"}:
            clients[key] = name
    book.close()
    return mapping, [{"key": key, "name": name} for key, name in sorted(clients.items(), key=lambda item: item[1])]


def main() -> None:
    alias_map, clients = aliases()
    daily: dict[tuple[str, date], int] = defaultdict(int)
    book = load_workbook(SOURCE, read_only=True, data_only=True)
    rows = book["OverviewOrder"].iter_rows(values_only=True)
    headers = {clean(value): index for index, value in enumerate(next(rows))}
    for row in rows:
        client_alias = clean(row[headers["Client ID"]]).upper()
        identity = alias_map.get(client_alias, {"key": client_alias.removeprefix("BR")})
        release_date = parse_date(row[headers["RELEASED"]])
        release_time = parse_time(row[headers["RELEASED2"]])
        lines = parse_lines(row[headers["NUM_LINES"]])
        if release_date and release_time and lines > 0:
            business_day = release_business_day(datetime.combine(release_date, release_time))
            daily[(identity["key"], business_day)] += lines
    book.close()

    daily_by_client: dict[str, list[int]] = defaultdict(list)
    for (client_key, _), lines in daily.items():
        daily_by_client[client_key].append(lines)

    for client in clients:
        volumes = daily_by_client.get(client["key"], [])
        p95 = percentile(volumes, .95)
        client["suggestedCapacity"] = int(math.ceil(p95)) if p95 is not None else None
        client["observedDays"] = len(volumes)
        client["averageDailyLines"] = round(sum(volumes) / len(volumes), 1) if volumes else None

    payload = {"clients": clients, "generatedAt": datetime.now().isoformat(timespec="seconds")}
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Gerado: {OUTPUT}")
    print(f"Clientes outbound: {len(clients)} | Dias cliente: {len(daily)}")


if __name__ == "__main__":
    main()
