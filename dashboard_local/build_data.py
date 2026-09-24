"""Gera os agregados consumidos pelo dashboard local de performance."""

from __future__ import annotations

import json
import math
import statistics
import csv
import os
from collections import defaultdict
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "00 - Bases OVERVIEW HTMLs.xlsx"
SLA_CONFIG = ROOT / "MetasSLA.xlsx"
OPERATIONAL_CONFIG = ROOT / "outputs" / "01a08633-89b4-7283-8814-f0b65f6970f7" / "ParametrosOperacionais_Inbound.xlsx"
OUTPUT = Path(__file__).resolve().parent / "data" / "dashboard_data.json"
OUTBOUND_CSV_DIR = ROOT / "Base_Clientes" / "base_consolidade_outbound"
INBOUND_CSV_DIR = ROOT / "Base_Clientes" / "base_consolidade_inbound"
INVENTORY_CSV = ROOT / "Base_Clientes" / "base_consolidada_inventário" / "inventario.csv"

STAGES = [
    ("Creation", "CREATION_DATE", "CREATION_DATE2"),
    ("Released", "RELEASED", "RELEASED2"),
    ("Allocated", "ALLOCATED", "ALLOCATED2"),
    ("In Progress", "IN_PROGRESS", "IN_PROGRESS2"),
    ("Picked", "PICKED", "PICKED2"),
    ("Packed", "PACKED", "PACKED2"),
    ("Ready to Load", "READY_TO_LOAD", "READY_TO_LOAD2"),
    ("Complete", "COMPLETE", "COMPLETE2"),
    ("Shipped", "SHIPPED", "SHIPPED2"),
]
STAGE_ORDER = {name: index for index, (name, _, _) in enumerate(STAGES)}
TRANSITIONS = [(STAGES[i][0], STAGES[i + 1][0]) for i in range(len(STAGES) - 1)] + [("Creation", "Shipped")]
SOURCE_AS_OF = datetime.fromtimestamp(SOURCE.stat().st_mtime) if SOURCE.exists() else datetime.now()
CSV_STAGE_COLUMNS = {
    "Creation": "Creation Date",
    "Released": "Released",
    "Allocated": "Allocated",
    "In Progress": "In Progress",
    "Picked": "Picked",
    "Packed": "Packed",
    "Ready to Load": "Ready To Load",
    "Complete": "Complete",
    "Shipped": "Shipped",
}
CSV_REQUIRED_COLUMNS = {"Site", "Client", "Owner", "Order Id", "Status", "Order Type", "Lines", *CSV_STAGE_COLUMNS.values()}
INBOUND_STAGES = ["Creation", "Released", "In Progress", "Finish", "UDF 1", "UDF 2", "UDF 3", "UDF 4"]
INBOUND_STAGE_ORDER = {name: index for index, name in enumerate(INBOUND_STAGES)}
INBOUND_TRANSITIONS = [("Creation", "Released"), ("Released", "In Progress"), ("In Progress", "Finish"), ("Creation", "Finish")]
INBOUND_CSV_STAGE_COLUMNS = {
    "Creation": "Creation Date",
    "Released": "Hold to Released",
    "In Progress": "Released to In progress",
    "Finish": "Finish Date",
    "UDF 1": "UDF_DATE 1",
    "UDF 2": "UDF_DATE 2",
    "UDF 3": "UDF_DATE 3",
    "UDF 4": "UDF_DATE 4",
}
INBOUND_REQUIRED_COLUMNS = {"Site", "Client", "Owner", "Pre-Advice ID", "Lines", "Status", "Type", "Due Date", *INBOUND_CSV_STAGE_COLUMNS.values()}


def clean(value: Any) -> str:
    return "" if value is None else str(value).strip()


def yes(value: Any) -> bool:
    return clean(value).lower() in {"sim", "s", "yes", "true", "1", "ativo"}


def number(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        parsed = float(str(value).replace(".", "").replace(",", ".")) if isinstance(value, str) else float(value)
        return max(0, int(round(parsed)))
    except (TypeError, ValueError):
        return None


def decimal(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None


def as_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        try:
            converted = from_excel(value)
            return converted.date() if isinstance(converted, datetime) else converted
        except (TypeError, ValueError, OverflowError):
            return None
    raw = clean(value)
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%y"):
        try:
            return datetime.strptime(raw[:10], fmt).date()
        except ValueError:
            pass
    return None


def as_time(value: Any) -> time | None:
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    if isinstance(value, (int, float)):
        seconds = round(float(value) * 86400) % 86400
        return time(seconds // 3600, seconds % 3600 // 60, seconds % 60)
    raw = clean(value)
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(raw, fmt).time()
        except ValueError:
            pass
    return None


def combined(date_value: Any, time_value: Any) -> datetime | None:
    day, clock = as_date(date_value), as_time(time_value)
    return datetime.combine(day, clock) if day and clock else None


def csv_timestamp(value: Any) -> datetime | None:
    raw = clean(value)
    if not raw:
        return None
    try:
        return datetime.strptime(raw, "%m-%d-%Y %H:%M:%S")
    except ValueError as exc:
        raise ValueError(f"timestamp inválido '{raw}'; esperado MM-DD-AAAA HH:MM:SS") from exc


def is_cancelled(value: Any) -> bool:
    """Pedidos só são cancelados para a análise quando CANCELLED contém uma data válida."""
    return as_date(value) is not None


def performance(value: Any) -> str:
    raw = clean(value).lower()
    if raw in {"on time", "ontime", "on-time"}:
        return "On Time"
    if raw in {"delay", "delayed", "late"}:
        return "Delay"
    return "Unclassified"


def quantile(values: list[float], q: float) -> float | None:
    if not values:
        return None
    ordered = sorted(values)
    position = (len(ordered) - 1) * q
    lo, hi = math.floor(position), math.ceil(position)
    if lo == hi:
        return ordered[lo]
    return ordered[lo] + (ordered[hi] - ordered[lo]) * (position - lo)


def table_rows(book, sheet_name: str) -> Iterable[dict[str, Any]]:
    if sheet_name not in book.sheetnames:
        return []
    raw_rows = list(book[sheet_name].iter_rows(values_only=True))
    header_index = next((i for i, row in enumerate(raw_rows) if sum(bool(clean(value)) for value in row) >= 2), None)
    if header_index is None:
        return []
    headers = [clean(value) for value in raw_rows[header_index]]
    return [dict(zip(headers, row)) for row in raw_rows[header_index + 1:] if any(value is not None and clean(value) for value in row)]


def load_sla_configuration() -> tuple[dict[str, dict[str, str]], dict[tuple[str, str], dict[str, float]]]:
    aliases: dict[str, dict[str, str]] = {}
    targets: dict[tuple[str, str], dict[str, float]] = {}
    if not SLA_CONFIG.exists():
        return aliases, targets
    book = load_workbook(SLA_CONFIG, read_only=True, data_only=True)
    for row in table_rows(book, "ClientAliases"):
        alias = clean(row.get("ClientAlias")).upper()
        if alias:
            aliases[alias] = {"key": clean(row.get("ClientKey")).upper(), "name": clean(row.get("ClientName")) or alias}
    for row in table_rows(book, "MetasSLA"):
        key, operation = clean(row.get("ClientKey")).upper(), clean(row.get("Operation"))
        if key and operation:
            targets[(key, operation)] = {"target": float(row.get("TargetPct") or .95), "warning": float(row.get("WarningPct") or .90)}
    book.close()
    return aliases, targets


class OperationalConfiguration:
    def __init__(self) -> None:
        self.capacities: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
        self.capacity_display: dict[tuple[str, str], dict[str, Any]] = {}
        self.inbound_capacities: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
        self.inbound_capacity_display: dict[tuple[str, str], dict[str, Any]] = {}
        self.journeys: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
        self.day_overrides: dict[str, dict[date, bool]] = defaultdict(dict)
        self.stage_targets: dict[tuple[str, str, str, str], list[dict[str, Any]]] = defaultdict(list)
        self.process_rules: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
        self.gross_rules: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
        self.inbound_rules: dict[tuple[str, str, str], list[dict[str, Any]]] = defaultdict(list)
        self.outbound_sources: dict[str, dict[str, Any]] = {}
        self.inbound_sources: dict[str, dict[str, Any]] = {}
        self.owner_aliases: dict[tuple[str, str], dict[str, str]] = {}
        self.rules = {"CoberturaConfiavel": .90, "CoberturaAtencao": .70, "AlertaDuracaoZero": .20, "HorasDiaUtil": 8.0}
        if not OPERATIONAL_CONFIG.exists():
            return
        book = load_workbook(OPERATIONAL_CONFIG, read_only=True, data_only=False)
        for row in table_rows(book, "Capacidade"):
            key = clean(row.get("ClientKey")).upper()
            owner = clean(row.get("OwnerKey")).upper() or "DEFAULT"
            official, suggested = number(row.get("CapacidadeOficial")), number(row.get("CapacidadeSugerida"))
            entry = {"value": official, "from": as_date(row.get("VigenciaInicio")) or date.min, "to": as_date(row.get("VigenciaFim")) or date.max, "active": yes(row.get("Ativo")) and official is not None and official > 0}
            if key:
                self.capacities[(key, owner)].append(entry)
                self.capacity_display[(key, owner)] = {"clientKey": key, "ownerKey": owner, "officialCapacity": official, "suggestedCapacity": suggested, "active": entry["active"], "validFrom": None if entry["from"] == date.min else entry["from"].isoformat(), "validTo": None if entry["to"] == date.max else entry["to"].isoformat()}
        for row in table_rows(book, "CapacidadeInbound"):
            key = clean(row.get("ClientKey")).upper()
            owner = clean(row.get("OwnerKey")).upper() or "DEFAULT"
            official, suggested = number(row.get("CapacidadeOficial")), number(row.get("CapacidadeSugerida"))
            entry = {"value": official, "from": as_date(row.get("VigenciaInicio")) or date.min, "to": as_date(row.get("VigenciaFim")) or date.max, "active": yes(row.get("Ativo")) and official is not None and official > 0}
            if key:
                self.inbound_capacities[(key, owner)].append(entry)
                self.inbound_capacity_display[(key, owner)] = {"clientKey": key, "ownerKey": owner, "officialCapacity": official, "suggestedCapacity": suggested, "active": entry["active"], "validFrom": None if entry["from"] == date.min else entry["from"].isoformat(), "validTo": None if entry["to"] == date.max else entry["to"].isoformat()}
        for row in table_rows(book, "Jornada"):
            key = clean(row.get("ClientKey")).upper()
            owner = clean(row.get("OwnerKey")).upper() or "DEFAULT"
            if key and yes(row.get("Ativo")):
                self.journeys[(key, owner)].append({"from": as_date(row.get("VigenciaInicio")) or date.min, "to": as_date(row.get("VigenciaFim")) or date.max, "segments": [(as_time(row.get("InicioManha")) or time(8), as_time(row.get("FimManha")) or time(12)), (as_time(row.get("InicioTarde")) or time(13), as_time(row.get("FimTarde")) or time(17))]})
        for row in table_rows(book, "Feriados"):
            day, key = as_date(row.get("Data")), clean(row.get("ClientKey")).upper() or "DEFAULT"
            if day and yes(row.get("Ativo")):
                self.day_overrides[key][day] = yes(row.get("DiaUtil"))
        for row in table_rows(book, "MetasEtapas"):
            key = clean(row.get("ClientKey")).upper() or "DEFAULT"
            owner = clean(row.get("OwnerKey")).upper() or "DEFAULT"
            start, end = clean(row.get("EtapaOrigem")), clean(row.get("EtapaDestino"))
            target = decimal(row.get("MetaHorasUteis"))
            if start and end and yes(row.get("Ativo")) and target is not None:
                self.stage_targets[(key, owner, start, end)].append({"value": target, "from": as_date(row.get("VigenciaInicio")) or date.min, "to": as_date(row.get("VigenciaFim")) or date.max})
        for row in table_rows(book, "ProcessosOutbound"):
            key = clean(row.get("ClientKey")).upper() or "DEFAULT"
            owner = clean(row.get("OwnerKey")).upper() or "DEFAULT"
            process_key = clean(row.get("ProcessoKey")).upper()
            process_name = clean(row.get("ProcessoNome")) or process_key
            start_from, start_to = clean(row.get("InicioOrigem")), clean(row.get("InicioDestino"))
            end_from, end_to = clean(row.get("FimOrigem")), clean(row.get("FimDestino"))
            if process_key and start_from and start_to and end_from and end_to and yes(row.get("Ativo")):
                self.process_rules[(key, owner, process_key)].append({
                    "processKey": process_key, "processName": process_name,
                    "startFrom": start_from, "startTo": start_to, "endFrom": end_from, "endTo": end_to,
                    "targetHours": decimal(row.get("MetaHorasUteis")),
                    "from": as_date(row.get("VigenciaInicio")) or date.min,
                    "to": as_date(row.get("VigenciaFim")) or date.max,
                    "ownerKey": owner,
                })
        for row in table_rows(book, "RegrasGross"):
            key = clean(row.get("ClientKey")).upper() or "DEFAULT"
            owner = clean(row.get("OwnerKey")).upper() or "DEFAULT"
            rule_id = clean(row.get("RegraID")).upper()
            start, end = clean(row.get("EtapaInicial")), clean(row.get("EtapaFinal"))
            if key and rule_id and yes(row.get("Ativo")):
                self.gross_rules[(key, owner, rule_id)].append({
                    "ruleId": rule_id,
                    "start": start,
                    "end": end,
                    "bandStartExclusive": as_time(row.get("HoraInicioExclusiva")),
                    "bandEndInclusive": as_time(row.get("HoraFimInclusiva")),
                    "deadlineType": clean(row.get("TipoPrazo")),
                    "businessDays": number(row.get("DiasUteis")),
                    "deadlineTime": as_time(row.get("HoraLimite")),
                    "businessHours": decimal(row.get("HorasUteis")),
                    "toleranceMinutes": number(row.get("ToleranciaMinutos")) or 0,
                    "from": as_date(row.get("VigenciaInicio")) or date.min,
                    "to": as_date(row.get("VigenciaFim")) or date.max,
                })
        for row in table_rows(book, "RegrasInbound"):
            key = clean(row.get("ClientKey")).upper() or "DEFAULT"
            owner = clean(row.get("OwnerKey")).upper() or "DEFAULT"
            rule_id = clean(row.get("RegraID")).upper()
            start, end = clean(row.get("EtapaInicial")), clean(row.get("EtapaFinal"))
            if key and rule_id and yes(row.get("Ativo")):
                self.inbound_rules[(key, owner, rule_id)].append({
                    "ruleId": rule_id,
                    "start": start,
                    "end": end,
                    "bandStartExclusive": as_time(row.get("HoraInicioExclusiva")),
                    "bandEndInclusive": as_time(row.get("HoraFimInclusiva")),
                    "deadlineType": clean(row.get("TipoPrazo")),
                    "businessDays": number(row.get("DiasUteis")),
                    "deadlineTime": as_time(row.get("HoraLimite")),
                    "businessHours": decimal(row.get("HorasUteis")),
                    "toleranceMinutes": number(row.get("ToleranciaMinutos")) or 0,
                    "from": as_date(row.get("VigenciaInicio")) or date.min,
                    "to": as_date(row.get("VigenciaFim")) or date.max,
                })
        for row in table_rows(book, "FontesOutbound"):
            key = clean(row.get("ClientKey")).upper()
            if not key or not yes(row.get("Ativo")):
                continue
            if key in self.outbound_sources:
                raise ValueError(f"mais de uma fonte Outbound ativa para {key}")
            cutoff = as_date(row.get("DataCorte"))
            filename = clean(row.get("ArquivoCSV"))
            if not cutoff or not filename:
                raise ValueError(f"fonte Outbound incompleta para {key}")
            self.outbound_sources[key] = {"clientKey": key, "filename": filename, "cutoff": cutoff, "sourceAlias": clean(row.get("ClientAliasFonte")).upper(), "ownerRequired": yes(row.get("OwnerObrigatorio"))}
        for row in table_rows(book, "FontesInbound"):
            key = clean(row.get("ClientKey")).upper()
            if not key or not yes(row.get("Ativo")):
                continue
            if key in self.inbound_sources:
                raise ValueError(f"mais de uma fonte Inbound ativa para {key}")
            cutoff = as_date(row.get("DataCorte"))
            filename = clean(row.get("ArquivoCSV"))
            if not cutoff or not filename:
                raise ValueError(f"fonte Inbound incompleta para {key}")
            self.inbound_sources[key] = {"clientKey": key, "filename": filename, "cutoff": cutoff, "sourceAlias": clean(row.get("ClientAliasFonte")).upper(), "ownerRequired": yes(row.get("OwnerObrigatorio"))}
        for row in table_rows(book, "OwnerAliases"):
            key = clean(row.get("ClientKey")).upper()
            alias = clean(row.get("OwnerAlias")).upper()
            if key and alias and yes(row.get("Ativo")):
                self.owner_aliases[(key, alias)] = {"key": clean(row.get("OwnerKey")).upper() or "NÃO INFORMADO", "name": clean(row.get("OwnerName")) or alias}
        for row in table_rows(book, "Regras"):
            key, value = clean(row.get("Parametro")), decimal(row.get("Valor"))
            if key in self.rules and value is not None:
                self.rules[key] = value
        book.close()

    @staticmethod
    def _effective(entries: list[dict[str, Any]], day: date) -> dict[str, Any] | None:
        valid = [entry for entry in entries if entry["from"] <= day <= entry["to"]]
        return max(valid, key=lambda item: item["from"]) if valid else None

    @staticmethod
    def _priority(client: str, owner: str) -> list[tuple[str, str]]:
        pairs = [(client, owner), (client, "DEFAULT"), ("DEFAULT", "DEFAULT")]
        return list(dict.fromkeys(pairs))

    def capacity(self, client: str, day: date, owner: str = "DEFAULT") -> int | None:
        entry = next((candidate for pair in self._priority(client, owner) if (candidate := self._effective(self.capacities.get(pair, []), day))), None)
        return entry["value"] if entry and entry["active"] else None

    def inbound_capacity(self, client: str, day: date, owner: str = "DEFAULT") -> int | None:
        entry = next((candidate for pair in self._priority(client, owner) if (candidate := self._effective(self.inbound_capacities.get(pair, []), day))), None)
        return entry["value"] if entry and entry["active"] else None

    def journey(self, client: str, day: date, owner: str = "DEFAULT") -> list[tuple[time, time]]:
        entry = next((candidate for pair in self._priority(client, owner) if (candidate := self._effective(self.journeys.get(pair, []), day))), None)
        return entry["segments"] if entry else [(time(8), time(12)), (time(13), time(17))]

    def is_business_day(self, client: str, day: date) -> bool:
        working = day.weekday() < 5
        if day in self.day_overrides.get("DEFAULT", {}):
            working = self.day_overrides["DEFAULT"][day]
        if day in self.day_overrides.get(client, {}):
            working = self.day_overrides[client][day]
        return working

    def stage_target(self, client: str, start: str, end: str, day: date, owner: str = "DEFAULT") -> float | None:
        entries = next((rows for c, o in self._priority(client, owner) if (rows := self.stage_targets.get((c, o, start, end), []))), [])
        entry = self._effective(entries, day)
        return entry["value"] if entry else None

    def process_rules_for(self, client: str, day: date, owner: str = "DEFAULT") -> list[dict[str, Any]]:
        selected: dict[str, dict[str, Any]] = {}
        for candidate_client, candidate_owner in self._priority(client, owner):
            for (rule_client, rule_owner, process_key), entries in self.process_rules.items():
                if (rule_client, rule_owner) != (candidate_client, candidate_owner) or process_key in selected:
                    continue
                entry = self._effective(entries, day)
                if entry:
                    selected[process_key] = entry
        return list(selected.values())

    def effective_gross_rule(self, client: str, stamps: dict[str, datetime | None], reference_day: date, owner: str = "DEFAULT") -> tuple[list[dict[str, Any]] | None, str | None]:
        return self._effective_rule(self.gross_rules, "Gross", client, stamps, reference_day, owner)

    def effective_inbound_rule(self, client: str, stamps: dict[str, datetime | None], reference_day: date, owner: str = "DEFAULT") -> tuple[list[dict[str, Any]] | None, str | None]:
        return self._effective_rule(self.inbound_rules, "Inbound", client, stamps, reference_day, owner)

    def _effective_rule(self, rules: dict[tuple[str, str, str], list[dict[str, Any]]], label: str, client: str, stamps: dict[str, datetime | None], reference_day: date, owner: str = "DEFAULT") -> tuple[list[dict[str, Any]] | None, str | None]:
        groups: list[tuple[str, list[dict[str, Any]]]] = []
        for candidate_client, candidate_owner in self._priority(client, owner):
            grouped: dict[str, list[dict[str, Any]]] = {}
            for raw_key, rows in rules.items():
                if len(raw_key) == 3:
                    key, own, rule_id = raw_key
                else:  # compatibilidade com configurações/testes anteriores ao Owner
                    key, rule_id = raw_key
                    own = "DEFAULT"
                if key == candidate_client and own == candidate_owner:
                    grouped[rule_id] = rows
            groups = list(grouped.items())
            if groups:
                break
        applicable: list[tuple[date, str, list[dict[str, Any]]]] = []
        for rule_id, rows in groups:
            start_name = rows[0]["start"] if rows else ""
            start_stamp = stamps.get(start_name)
            rule_day = start_stamp.date() if start_stamp else reference_day
            active_rows = [row for row in rows if row["from"] <= rule_day <= row["to"]]
            if active_rows:
                applicable.append((max(row["from"] for row in active_rows), rule_id, active_rows))
        if not applicable:
            return None, "Nenhuma regra ativa para o cliente e a vigência"
        newest = max(item[0] for item in applicable)
        candidates = [item for item in applicable if item[0] == newest]
        if len(candidates) != 1:
            return None, f"Mais de uma regra {label} ativa na mesma vigência"
        return candidates[0][2], None

    def gross_rule_summary(self, client: str, owner: str = "DEFAULT") -> dict[str, Any] | None:
        return self._rule_summary(self.gross_rules, client, owner)

    def inbound_rule_summary(self, client: str, owner: str = "DEFAULT") -> dict[str, Any] | None:
        return self._rule_summary(self.inbound_rules, client, owner)

    def _rule_summary(self, rules: dict[tuple[str, str, str], list[dict[str, Any]]], client: str, owner: str = "DEFAULT") -> dict[str, Any] | None:
        groups: list[tuple[str, list[dict[str, Any]]]] = []
        inherited = True
        for candidate_client, candidate_owner in self._priority(client, owner):
            grouped: dict[str, list[dict[str, Any]]] = {}
            for raw_key, rows in rules.items():
                if len(raw_key) == 3:
                    key, own, rule_id = raw_key
                else:
                    key, rule_id = raw_key
                    own = "DEFAULT"
                if key == candidate_client and own == candidate_owner:
                    grouped[rule_id] = rows
            groups = list(grouped.items())
            if groups:
                inherited = (candidate_client, candidate_owner) != (client, owner)
                break
        if not groups:
            return None
        rule_id, rows = max(groups, key=lambda item: max(row["from"] for row in item[1]))
        first = rows[0]
        return {"ruleId": rule_id, "stageFrom": first["start"], "stageTo": first["end"], "bands": len(rows), "inherited": inherited}

    def owner_identity(self, client: str, raw_owner: Any) -> dict[str, str]:
        alias = clean(raw_owner).upper()
        if not alias:
            return {"key": "NÃO INFORMADO", "name": "Não informado", "raw": ""}
        mapped = self.owner_aliases.get((client, alias))
        if mapped:
            return {**mapped, "raw": clean(raw_owner)}
        return {"key": alias, "name": alias, "raw": clean(raw_owner)}

    def capacity_profile(self, client: str, owner: str) -> dict[str, Any]:
        for pair in self._priority(client, owner):
            if pair in self.capacity_display:
                return {**self.capacity_display[pair], "operation": "Outbound", "appliedToOwner": owner, "inherited": pair != (client, owner)}
        return {"clientKey": client, "ownerKey": owner, "operation": "Outbound", "officialCapacity": None, "suggestedCapacity": None, "active": False, "validFrom": None, "validTo": None, "appliedToOwner": owner, "inherited": True}

    def inbound_capacity_profile(self, client: str, owner: str) -> dict[str, Any]:
        for pair in self._priority(client, owner):
            if pair in self.inbound_capacity_display:
                return {**self.inbound_capacity_display[pair], "operation": "Inbound", "appliedToOwner": owner, "inherited": pair != (client, owner)}
        return {"clientKey": client, "ownerKey": owner, "operation": "Inbound", "officialCapacity": None, "suggestedCapacity": None, "active": False, "validFrom": None, "validTo": None, "appliedToOwner": owner, "inherited": True}


class BusinessCalendar:
    START = date(2024, 1, 1)
    END = date(2028, 12, 31)

    def __init__(self, config: OperationalConfiguration) -> None:
        self.config = config
        self.cache: dict[str, dict[str, Any]] = {}

    def _client(self, client: str, owner: str = "DEFAULT") -> dict[str, Any]:
        cache_key = f"{client}|{owner}"
        if cache_key in self.cache:
            return self.cache[cache_key]
        dates = [self.START + timedelta(days=i) for i in range((self.END - self.START).days + 1)]
        segments = {day: self.config.journey(client, day, owner) if self.config.is_business_day(client, day) else [] for day in dates}
        totals = [sum((datetime.combine(day, end) - datetime.combine(day, start)).total_seconds() for start, end in segments[day]) for day in dates]
        prefix = [0.0]
        for value in totals:
            prefix.append(prefix[-1] + value)
        data = {"index": {day: i for i, day in enumerate(dates)}, "segments": segments, "prefix": prefix}
        self.cache[cache_key] = data
        return data

    def normalize_release(self, client: str, stamp: datetime, owner: str = "DEFAULT") -> datetime:
        data = self._client(client, owner)
        day = max(self.START, min(self.END, stamp.date()))
        while True:
            segments = data["segments"].get(day, [])
            if not segments:
                day += timedelta(days=1)
                continue
            for start, end in segments:
                if stamp.date() < day or stamp.time() < start:
                    return datetime.combine(day, start)
                if start <= stamp.time() < end:
                    return stamp
            day += timedelta(days=1)

    def add_business_days(self, client: str, day: date, amount: int, owner: str = "DEFAULT") -> date:
        current = day
        remaining = max(0, amount)
        while remaining:
            current += timedelta(days=1)
            if self.config.is_business_day(client, current):
                remaining -= 1
        return current

    def add_business_hours(self, client: str, stamp: datetime, amount: float, owner: str = "DEFAULT") -> datetime:
        current = self.normalize_release(client, stamp, owner)
        remaining = max(0.0, amount) * 3600
        while remaining > 0:
            segments = self._client(client, owner)["segments"].get(current.date(), [])
            progressed = False
            for segment_start, segment_end in segments:
                start = max(current, datetime.combine(current.date(), segment_start))
                end = datetime.combine(current.date(), segment_end)
                if start >= end:
                    continue
                available = (end - start).total_seconds()
                if remaining <= available:
                    return start + timedelta(seconds=remaining)
                remaining -= available
                current = end
                progressed = True
            current = self.normalize_release(client, datetime.combine(current.date() + timedelta(days=1), time.min), owner)
            if not progressed and not segments:
                continue
        return current

    @staticmethod
    def _overlap(day: date, segments: list[tuple[time, time]], start: datetime, end: datetime) -> float:
        total = 0.0
        for segment_start, segment_end in segments:
            left = max(start, datetime.combine(day, segment_start))
            right = min(end, datetime.combine(day, segment_end))
            if right > left:
                total += (right - left).total_seconds()
        return total

    def business_hours(self, client: str, start: datetime, end: datetime, owner: str = "DEFAULT") -> float:
        if end <= start:
            return 0.0
        data = self._client(client, owner)
        if start.date() == end.date():
            return self._overlap(start.date(), data["segments"].get(start.date(), []), start, end) / 3600
        first = self._overlap(start.date(), data["segments"].get(start.date(), []), start, datetime.combine(start.date(), time.max))
        last = self._overlap(end.date(), data["segments"].get(end.date(), []), datetime.combine(end.date(), time.min), end)
        start_index, end_index = data["index"].get(start.date()), data["index"].get(end.date())
        middle = 0.0 if start_index is None or end_index is None else data["prefix"][end_index] - data["prefix"][start_index + 1]
        return max(0.0, first + middle + last) / 3600


def identity(value: Any, aliases: dict[str, dict[str, str]]) -> dict[str, str]:
    alias = clean(value).upper() or "SEM_CLIENTE"
    return aliases.get(alias, {"key": alias.removeprefix("BR"), "name": alias.removeprefix("BR")})


def inventory_timestamp(value: Any) -> datetime | None:
    raw = clean(value)
    if not raw:
        return None
    try:
        return datetime.fromisoformat(raw)
    except ValueError:
        pass
    for fmt in ("%d-%m-%Y %H:%M", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(raw, fmt)
        except ValueError:
            pass
    return None


def load_inventory(aliases: dict[str, dict[str, str]], operational: OperationalConfiguration) -> dict[str, Any]:
    """Read the partial stock-check history and derive the latest known line per key."""
    empty = {"meta": {"source": str(INVENTORY_CSV.relative_to(ROOT)), "rows": 0, "identifiedRows": 0, "noInventoryRows": 0, "lists": 0, "minCountDate": None, "maxCountDate": None}, "latestRows": [], "history": [], "summaries": [], "quality": {}}
    if not INVENTORY_CSV.exists():
        return empty
    with INVENTORY_CSV.open("r", encoding="utf-8-sig", newline="") as handle:
        raw_rows = list(csv.reader(handle))
    if len(raw_rows) < 3:
        return empty
    source_headers = raw_rows[2]
    seen_headers: dict[str, int] = {}
    headers: list[str] = []
    for index, header in enumerate(source_headers):
        base = clean(header) or f"blank_{index + 1}"
        seen_headers[base] = seen_headers.get(base, 0) + 1
        headers.append(base if seen_headers[base] == 1 else f"{base}_{seen_headers[base]}")
    # The export leaves the final four headers on the preceding visual row.
    headers[18:22] = ["Final Qty", "Final Difference", "Final Value (BRL)", "Value of Discrepancy (BRL)"]
    latest: dict[tuple[str, str, str, str, str], dict[str, Any]] = {}
    history_detail_rows: list[dict[str, Any]] = []
    history: dict[tuple[str, str, str], dict[str, Any]] = {}
    lists: set[str] = set()
    all_dates: list[datetime] = []
    quality = {"rows": 0, "identifiedRows": 0, "noInventoryRows": 0, "blankClientRows": 0, "lists": 0, "rowsWithSecondCount": 0, "rowsWithDifference": 0, "foundRows": 0, "lostRows": 0, "okRows": 0, "repeatedLineKeys": 0, "missingCountDates": 0}
    historical_keys: set[tuple[str, str, str, str, str]] = set()
    repeated_keys: set[tuple[str, str, str, str, str]] = set()
    for raw in raw_rows[3:]:
        values = (raw + [""] * len(headers))[:len(headers)]
        row = dict(zip(headers, values))
        quality["rows"] += 1
        list_id = clean(row.get("List ID"))
        lists.add(list_id)
        client_raw = clean(row.get("Client ID")).upper()
        person = identity(client_raw, aliases) if client_raw else {"key": "SEM_CLIENTE", "name": "Sem cliente"}
        owner_raw = clean(row.get("Owner ID")).upper()
        owner = operational.owner_identity(person["key"], owner_raw) if client_raw else {"key": "SEM_OWNER", "name": "Sem owner", "raw": owner_raw}
        sku = clean(row.get("SKU"))
        no_inventory = not client_raw or sku.upper() == "NOINVENTORY"
        if no_inventory:
            quality["noInventoryRows"] += 1
        else:
            quality["identifiedRows"] += 1
        if not client_raw:
            quality["blankClientRows"] += 1
        first_date = inventory_timestamp(row.get("1st Count Date"))
        second_date = inventory_timestamp(row.get("2nd Count Date"))
        count_date = second_date or first_date
        if count_date:
            all_dates.append(count_date)
        else:
            quality["missingCountDates"] += 1
        if second_date:
            quality["rowsWithSecondCount"] += 1
        type_name = clean(row.get("Type")) or "OK"
        quality[f"{type_name.lower()}Rows"] = quality.get(f"{type_name.lower()}Rows", 0) + 1
        original_qty = number(row.get("Original Quantity"))
        first_qty = number(row.get("1st Count"))
        second_qty = number(row.get("2nd Count"))
        final_qty = number(row.get("Final Qty"))
        if final_qty is None:
            final_qty = second_qty if second_qty is not None else first_qty
        difference = number(row.get("Final Difference"))
        if difference is None and original_qty is not None and final_qty is not None:
            difference = final_qty - original_qty
        if difference not in (None, 0):
            quality["rowsWithDifference"] += 1
        item = {"listId": list_id, "clientKey": person["key"], "clientName": person["name"], "ownerKey": owner["key"], "ownerName": owner["name"], "ownerRaw": owner_raw, "location": clean(row.get("From Location")), "tagId": clean(row.get("Tag ID")), "sku": sku, "description": clean(row.get("Description")), "originalQty": original_qty, "firstCount": first_qty, "secondCount": second_qty, "finalQty": final_qty, "difference": difference, "finalValue": decimal(row.get("Final Value (BRL)")), "discrepancyValue": decimal(row.get("Value of Discrepancy (BRL)")), "type": type_name, "firstCountDate": first_date.isoformat(timespec="minutes") if first_date else None, "secondCountDate": second_date.isoformat(timespec="minutes") if second_date else None, "countDate": count_date.isoformat(timespec="minutes") if count_date else None, "countDay": count_date.date().isoformat() if count_date else None, "countMonth": count_date.strftime("%Y-%m") if count_date else None, "noInventory": no_inventory}
        if no_inventory or not count_date:
            continue
        history_detail_rows.append(item)
        key = (person["key"], owner["key"], item["location"], item["tagId"], sku)
        if key in historical_keys:
            repeated_keys.add(key)
        historical_keys.add(key)
        current = latest.get(key)
        if current is None or (count_date, list_id) >= (inventory_timestamp(current.get("countDate")) or datetime.min, current.get("listId", "")):
            latest[key] = item
        period_key = (person["key"], owner["key"], item["countMonth"] or "Sem data")
        bucket = history.setdefault(period_key, {"clientKey": person["key"], "clientName": person["name"], "ownerKey": owner["key"], "ownerName": owner["name"], "month": period_key[2], "countedRows": 0, "countedQty": 0, "difference": 0, "absoluteDifference": 0, "lists": set()})
        bucket["countedRows"] += 1
        bucket["countedQty"] += final_qty or 0
        bucket["difference"] += difference or 0
        bucket["absoluteDifference"] += abs(difference or 0)
        bucket["lists"].add(list_id)
    quality["repeatedLineKeys"] = len(repeated_keys)
    quality["lists"] = len(lists)
    latest_rows = sorted(latest.values(), key=lambda item: (item["clientName"], item["ownerName"], item["location"], item["sku"], item["tagId"]))
    summary_map: dict[tuple[str, str], dict[str, Any]] = {}
    for item in latest_rows:
        key = (item["clientKey"], item["ownerKey"])
        summary = summary_map.setdefault(key, {"clientKey": item["clientKey"], "clientName": item["clientName"], "ownerKey": item["ownerKey"], "ownerName": item["ownerName"], "lineCount": 0, "quantity": 0, "originalQuantity": 0, "difference": 0, "absoluteDifference": 0, "linesWithDifference": 0, "foundLines": 0, "lostLines": 0, "lastCountDate": None})
        summary["lineCount"] += 1
        summary["quantity"] += item["finalQty"] or 0
        summary["originalQuantity"] += item["originalQty"] or 0
        summary["difference"] += item["difference"] or 0
        summary["absoluteDifference"] += abs(item["difference"] or 0)
        summary["linesWithDifference"] += int((item["difference"] or 0) != 0)
        summary["foundLines"] += int(item["type"] == "Found")
        summary["lostLines"] += int(item["type"] == "Lost")
        if item["countDate"] and (summary["lastCountDate"] is None or item["countDate"] > summary["lastCountDate"]):
            summary["lastCountDate"] = item["countDate"]
    history_summary_rows = [{**item, "lists": len(item["lists"])} for item in history.values()]
    history_summary_rows.sort(key=lambda item: (item["month"], item["clientName"], item["ownerName"]))
    summary_rows = sorted(summary_map.values(), key=lambda item: (item["clientName"], item["ownerName"]))
    return {"meta": {"source": str(INVENTORY_CSV.relative_to(ROOT)), "rows": quality["rows"], "identifiedRows": quality["identifiedRows"], "noInventoryRows": quality["noInventoryRows"], "lists": quality["lists"], "minCountDate": min(all_dates).isoformat(timespec="minutes") if all_dates else None, "maxCountDate": max(all_dates).isoformat(timespec="minutes") if all_dates else None}, "latestRows": latest_rows, "historyRows": history_detail_rows, "history": history_summary_rows, "summaries": summary_rows, "quality": quality}


def new_metric(key: str, operation: str, month: str, owner: str = "ALL", source_kind: str = "legacy") -> dict[str, Any]:
    return {"clientKey": key, "ownerKey": owner, "operation": operation, "month": month, "sourceKind": source_kind, "records": 0, "eligibleRecords": 0, "validLines": 0, "onTimeLines": 0, "delayedLines": 0, "unclassifiedRecords": 0, "missingLinesRecords": 0, "calcGrossEligibleRecords": 0, "calcGrossEligibleLines": 0, "calcGrossOnTimeLines": 0, "calcGrossDelayedLines": 0, "calcCompletedRecords": 0, "calcCompletedEligibleLines": 0, "calcCompletedOnTimeLines": 0, "calcCompletedDelayedLines": 0, "calcOpenWithinRecords": 0, "calcOpenWithinLines": 0, "calcOpenOverdueRecords": 0, "calcOpenOverdueLines": 0, "calcFallbackRecords": 0, "calcFallbackLines": 0, "calcNotCalculatedRecords": 0, "calcNotCalculatedLines": 0, "comparisonEligibleLines": 0, "comparisonBothOnTimeLines": 0, "comparisonBothDelayLines": 0, "comparisonSourceOnTimeCalcDelayLines": 0, "comparisonSourceDelayCalcOnTimeLines": 0, "netEligibleLines": 0, "netOnTimeLines": 0, "netDelayedLines": 0, "capacityExcludedLines": 0, "netUnavailableLines": 0, "openMetricsAvailableRecords": 0, "newSourceRecords": 0}


def new_stage(operation: str, key: str, month: str, start: str, end: str, owner: str = "HISTORICO") -> dict[str, Any]:
    return {
        "operation": operation, "clientKey": key, "ownerKey": owner, "month": month,
        "stageFrom": start, "stageTo": end, "totalRecords": 0, "validCount": 0,
        "missingCount": 0, "openCount": 0, "incompleteCount": 0, "invalidSequenceCount": 0,
        "zeroCount": 0, "calendarValues": [], "businessValues": [], "openAgeValues": [],
        # The process cards use the consolidated source and are weighted by lines.
        "performanceSourceRecords": 0, "performanceTotalRecords": 0, "performanceTotalLines": 0,
        "performanceEligibleRecords": 0, "performanceEligibleLines": 0,
        "performanceOnTimeLines": 0, "performanceDelayedLines": 0,
        "performanceOpenWithinLines": 0, "performanceNotCalculatedLines": 0,
        "performanceNoSlaRecords": 0,
    }


def new_process(operation: str, key: str, month: str, rule: dict[str, Any], owner: str = "DEFAULT") -> dict[str, Any]:
    return {
        "operation": operation, "clientKey": key, "ownerKey": owner, "month": month,
        "processKey": rule["processKey"], "processName": rule["processName"],
        "startFrom": rule["startFrom"], "startTo": rule["startTo"], "endFrom": rule["endFrom"], "endTo": rule["endTo"],
        "targetHours": rule.get("targetHours"), "totalRecords": 0, "totalLines": 0,
        "deadlineSource": "MetaHorasUteis", "deadlineType": None,
        "eligibleRecords": 0, "eligibleLines": 0, "onTimeLines": 0, "delayedLines": 0,
        "openWithinLines": 0, "notCalculatedLines": 0, "noSlaRecords": 0,
    }


def new_combined_process(operation: str, key: str, period: str, owner: str = "DEFAULT") -> dict[str, Any]:
    return {
        "operation": operation, "clientKey": key, "ownerKey": owner, "month": period,
        "processKey": "OPERACIONAL_COMBINADA", "processName": "Performance operacional combinada",
        "startFrom": "Allocated", "startTo": "In Progress", "endFrom": "Packed", "endTo": "Ready to Load",
        "targetHours": None, "totalRecords": 0, "totalLines": 0,
        "eligibleRecords": 0, "eligibleLines": 0, "onTimeLines": 0, "delayedLines": 0,
        "openWithinLines": 0, "notCalculatedLines": 0, "noSlaRecords": 0,
    }


def classify_process_performance(calendar: BusinessCalendar, client: str, stamps: dict[str, datetime | None], rule: dict[str, Any], as_of: datetime, owner: str = "DEFAULT", config: OperationalConfiguration | None = None, reference_day: date | None = None) -> dict[str, Any]:
    """Measure the configured process from the start transition origin to the end transition destination."""
    start_name, end_name = rule["startFrom"], rule["endTo"]
    if start_name not in STAGE_ORDER or end_name not in STAGE_ORDER or STAGE_ORDER[start_name] >= STAGE_ORDER[end_name]:
        return {"status": "Not Calculated", "performance": None, "reason": "Limites de processo inválidos"}
    start, end = stamps.get(start_name), stamps.get(end_name)
    target_hours = rule.get("targetHours")
    gross_rows = None
    if config and reference_day and start:
        gross_rows, _ = config.effective_gross_rule(client, {"Allocated": start}, reference_day, owner)
    if gross_rows:
        normalized_start = calendar.normalize_release(client, start, owner)
        clock = normalized_start.time()
        bands = [row for row in gross_rows if (row["bandStartExclusive"] is None or clock > row["bandStartExclusive"]) and (row["bandEndInclusive"] is None or clock <= row["bandEndInclusive"])]
        if len(bands) != 1:
            return {"status": "Not Calculated", "performance": None, "reason": "Faixa de corte Gross ausente ou sobreposta", "deadlineSource": "Gross"}
        band = bands[0]
        if band["deadlineType"] == "Dia útil + hora":
            due_day = calendar.add_business_days(client, normalized_start.date(), band["businessDays"], owner) if band["businessDays"] is not None else None
            due = datetime.combine(due_day, band["deadlineTime"]) if due_day and band["deadlineTime"] else None
        elif band["deadlineType"] == "Horas úteis":
            due = calendar.add_business_hours(client, normalized_start, band["businessHours"], owner) if band["businessHours"] is not None else None
        else:
            due = None
        if due is None:
            return {"status": "Not Calculated", "performance": None, "reason": "Prazo Gross incompleto ou inválido", "deadlineSource": "Gross"}
        due += timedelta(minutes=band["toleranceMinutes"])
        base = {"deadlineSource": "Gross", "deadlineType": band["deadlineType"], "dueAt": due.isoformat(timespec="seconds"), "startAt": start.isoformat(timespec="seconds")}
        if end:
            if end < start:
                return {**base, "status": "Not Calculated", "performance": None, "reason": "Evento final anterior ao evento inicial", "endAt": end.isoformat(timespec="seconds")}
            on_time = end <= due
            return {**base, "status": "Completed On Time" if on_time else "Completed Delay", "performance": "On Time" if on_time else "Delay", "duration": calendar.business_hours(client, start, end, owner), "endAt": end.isoformat(timespec="seconds")}
        later_exists = any(stamps.get(stage_name) for stage_name, _, _ in STAGES[STAGE_ORDER[end_name] + 1:])
        if later_exists:
            return {**base, "status": "Not Calculated", "performance": None, "reason": f"Evento final ausente apesar de etapa posterior: {end_name}"}
        return {**base, "status": "Open Delay" if as_of > due else "Open On Track", "performance": "Delay" if as_of > due else None}
    if target_hours is None:
        return {"status": "No SLA", "performance": None, "reason": "Meta de processo ausente"}
    if not start:
        return {"status": "Not Calculated", "performance": None, "reason": f"Evento inicial ausente: {start_name}"}
    if end:
        if end < start:
            return {"status": "Not Calculated", "performance": None, "reason": "Evento final anterior ao evento inicial"}
        duration = calendar.business_hours(client, start, end, owner)
        on_time = duration <= target_hours
        return {"status": "Completed On Time" if on_time else "Completed Delay", "performance": "On Time" if on_time else "Delay", "duration": duration}
    if any(stamps.get(stage_name) for stage_name, _, _ in STAGES[STAGE_ORDER[end_name] + 1:]):
        return {"status": "Not Calculated", "performance": None, "reason": f"Evento final ausente apesar de etapa posterior: {end_name}"}
    overdue = calendar.business_hours(client, start, as_of, owner) > target_hours
    return {"status": "Open Delay" if overdue else "Open On Track", "performance": "Delay" if overdue else None}


def classify_stage_performance(calendar: BusinessCalendar, client: str, stamps: dict[str, datetime | None], start_name: str, end_name: str, target_hours: float | None, as_of: datetime, owner: str = "DEFAULT") -> dict[str, Any]:
    """Classify one intermediate outbound process without applying endpoint fallback."""
    if target_hours is None:
        return {"status": "No SLA", "performance": None, "reason": "Meta de etapa ausente"}
    start = stamps.get(start_name)
    if not start:
        return {"status": "Not Calculated", "performance": None, "reason": f"Evento inicial ausente: {start_name}"}
    end = stamps.get(end_name)
    if end:
        if end < start:
            return {"status": "Not Calculated", "performance": None, "reason": "Evento final anterior ao evento inicial"}
        duration = calendar.business_hours(client, start, end, owner)
        on_time = duration <= target_hours
        return {"status": "Completed On Time" if on_time else "Completed Delay", "performance": "On Time" if on_time else "Delay", "duration": duration}
    end_position = STAGE_ORDER[end_name]
    if any(stamps.get(stage_name) for stage_name, _, _ in STAGES[end_position + 1:]):
        return {"status": "Not Calculated", "performance": None, "reason": f"Evento final ausente apesar de etapa posterior: {end_name}"}
    overdue = calendar.business_hours(client, start, as_of, owner) > target_hours
    return {"status": "Open Delay" if overdue else "Open On Track", "performance": "Delay" if overdue else None}


def split_capacity(lines: int, remaining: int) -> tuple[int, int, int]:
    """Divide as linhas do pedido entre NET e excesso, preservando FIFO."""
    included = min(lines, max(0, remaining))
    excluded = lines - included
    return included, excluded, remaining - included


def allocate_inbound_net_capacity(records: list[dict[str, Any]], official: int | None) -> list[tuple[int, int]]:
    """Retira do NET apenas o atraso que excede a capacidade diária."""
    if official is None:
        return [(0, 0) for _ in records]
    total_lines = sum(max(0, int(record.get("lines") or 0)) for record in records)
    excess = max(0, total_lines - official)
    allocations: list[tuple[int, int]] = []
    for record in records:
        lines = max(0, int(record.get("lines") or 0))
        excluded = min(lines, excess) if record.get("performance") == "Delay" else 0
        excess -= excluded
        allocations.append((lines - excluded, excluded))
    return allocations


def classify_gross(config: OperationalConfiguration, calendar: BusinessCalendar, client: str, stamps: dict[str, datetime | None], reference_day: date, as_of: datetime = SOURCE_AS_OF, owner: str = "DEFAULT") -> dict[str, Any]:
    rows, rule_error = config.effective_gross_rule(client, stamps, reference_day, owner)
    if not rows:
        return {"status": "Not Calculated", "performance": None, "reason": rule_error}
    start_name, end_name = rows[0]["start"], rows[0]["end"]
    if start_name not in STAGE_ORDER or end_name not in STAGE_ORDER or STAGE_ORDER[start_name] >= STAGE_ORDER[end_name]:
        return {"status": "Not Calculated", "performance": None, "reason": "Sequência de eventos inválida", "ruleId": rows[0]["ruleId"], "stageFrom": start_name, "stageTo": end_name}
    if any(row["start"] != start_name or row["end"] != end_name for row in rows):
        return {"status": "Not Calculated", "performance": None, "reason": "Eventos divergentes dentro da mesma RegraID", "ruleId": rows[0]["ruleId"], "stageFrom": start_name, "stageTo": end_name}
    start = stamps.get(start_name)
    if not start:
        return {"status": "Not Calculated", "performance": None, "reason": f"Evento inicial ausente: {start_name}", "ruleId": rows[0]["ruleId"], "stageFrom": start_name, "stageTo": end_name}
    normalized_start = calendar.normalize_release(client, start, owner)
    clock = normalized_start.time()
    bands = [row for row in rows if (row["bandStartExclusive"] is None or clock > row["bandStartExclusive"]) and (row["bandEndInclusive"] is None or clock <= row["bandEndInclusive"])]
    if len(bands) != 1:
        reason = "Nenhuma faixa cobre o horário inicial" if not bands else "Faixas de horário sobrepostas"
        return {"status": "Not Calculated", "performance": None, "reason": reason, "ruleId": rows[0]["ruleId"], "stageFrom": start_name, "stageTo": end_name, "startAt": start.isoformat(timespec="seconds")}
    rule = bands[0]
    if rule["deadlineType"] == "Dia útil + hora":
        if rule["businessDays"] is None or rule["deadlineTime"] is None:
            due = None
        else:
            due_day = calendar.add_business_days(client, normalized_start.date(), rule["businessDays"], owner)
            due = datetime.combine(due_day, rule["deadlineTime"])
    elif rule["deadlineType"] == "Horas úteis":
        due = calendar.add_business_hours(client, normalized_start, rule["businessHours"], owner) if rule["businessHours"] is not None else None
    else:
        due = None
    if due is None:
        return {"status": "Not Calculated", "performance": None, "reason": "Prazo incompleto ou tipo de prazo inválido", "ruleId": rule["ruleId"], "stageFrom": start_name, "stageTo": end_name, "startAt": start.isoformat(timespec="seconds")}
    due += timedelta(minutes=rule["toleranceMinutes"])
    end = stamps.get(end_name)
    end_name_used = end_name
    fallback_reason = None
    if not end:
        for candidate_name, _, _ in reversed(STAGES[STAGE_ORDER[start_name] + 1:STAGE_ORDER[end_name]]):
            candidate_stamp = stamps.get(candidate_name)
            if candidate_stamp:
                end = candidate_stamp
                end_name_used = candidate_name
                fallback_reason = f"{end_name} ausente; usada a etapa anterior {candidate_name}"
                break
    base = {"ruleId": rule["ruleId"], "stageFrom": start_name, "stageTo": end_name, "startAt": start.isoformat(timespec="seconds"), "dueAt": due.isoformat(timespec="seconds")}
    if end:
        if end < start:
            return {**base, "stageToUsed": end_name_used, "fallbackApplied": bool(fallback_reason), "fallbackReason": fallback_reason, "status": "Not Calculated", "performance": None, "reason": "Evento final anterior ao evento inicial", "endAt": end.isoformat(timespec="seconds")}
        on_time = end <= due
        return {**base, "stageToUsed": end_name_used, "fallbackApplied": bool(fallback_reason), "fallbackReason": fallback_reason, "status": "Completed On Time" if on_time else "Completed Delay", "performance": "On Time" if on_time else "Delay", "reason": None, "endAt": end.isoformat(timespec="seconds")}
    later_exists = any(stamps.get(name) for name, _, _ in STAGES[STAGE_ORDER[end_name] + 1:])
    if later_exists:
        return {**base, "status": "Not Calculated", "performance": None, "reason": f"Evento final ausente apesar de etapa posterior: {end_name}"}
    overdue = as_of > due
    return {**base, "status": "Open Delay" if overdue else "Open On Track", "performance": "Delay" if overdue else None, "reason": None}


def classify_inbound(config: OperationalConfiguration, calendar: BusinessCalendar, client: str, stamps: dict[str, datetime | None], source_due: datetime | None, reference_day: date, as_of: datetime = SOURCE_AS_OF, owner: str = "DEFAULT") -> dict[str, Any]:
    rows, rule_error = config.effective_inbound_rule(client, stamps, reference_day, owner)
    if not rows:
        return {"status": "Not Calculated", "performance": None, "reason": rule_error}
    start_name, end_name = rows[0]["start"], rows[0]["end"]
    if start_name not in INBOUND_STAGE_ORDER or end_name not in INBOUND_STAGE_ORDER or INBOUND_STAGE_ORDER[start_name] >= INBOUND_STAGE_ORDER[end_name]:
        return {"status": "Not Calculated", "performance": None, "reason": "Sequência de eventos inválida", "ruleId": rows[0]["ruleId"], "stageFrom": start_name, "stageTo": end_name}
    if any(row["start"] != start_name or row["end"] != end_name for row in rows):
        return {"status": "Not Calculated", "performance": None, "reason": "Eventos divergentes dentro da mesma RegraID", "ruleId": rows[0]["ruleId"], "stageFrom": start_name, "stageTo": end_name}
    start = stamps.get(start_name)
    if not start:
        return {"status": "Not Calculated", "performance": None, "reason": f"Evento inicial ausente: {start_name}", "ruleId": rows[0]["ruleId"], "stageFrom": start_name, "stageTo": end_name}
    normalized_start = calendar.normalize_release(client, start, owner)
    clock = normalized_start.time()
    bands = [row for row in rows if (row["bandStartExclusive"] is None or clock > row["bandStartExclusive"]) and (row["bandEndInclusive"] is None or clock <= row["bandEndInclusive"])]
    if len(bands) != 1:
        reason = "Nenhuma faixa cobre o horário inicial" if not bands else "Faixas de horário sobrepostas"
        return {"status": "Not Calculated", "performance": None, "reason": reason, "ruleId": rows[0]["ruleId"], "stageFrom": start_name, "stageTo": end_name, "startAt": start.isoformat(timespec="seconds")}
    rule = bands[0]
    if rule["deadlineType"] == "Data de vencimento da fonte":
        due = source_due
    elif rule["deadlineType"] == "Dia útil + hora":
        due_day = calendar.add_business_days(client, normalized_start.date(), rule["businessDays"], owner) if rule["businessDays"] is not None else None
        due = datetime.combine(due_day, rule["deadlineTime"]) if due_day and rule["deadlineTime"] else None
    elif rule["deadlineType"] == "Horas úteis":
        due = calendar.add_business_hours(client, normalized_start, rule["businessHours"], owner) if rule["businessHours"] is not None else None
    else:
        due = None
    if due is None:
        reason = "Due Date ausente" if rule["deadlineType"] == "Data de vencimento da fonte" else "Prazo incompleto ou tipo de prazo inválido"
        return {"status": "Not Calculated", "performance": None, "reason": reason, "ruleId": rule["ruleId"], "stageFrom": start_name, "stageTo": end_name, "startAt": start.isoformat(timespec="seconds")}
    due += timedelta(minutes=rule["toleranceMinutes"])
    end = stamps.get(end_name)
    end_name_used = end_name
    fallback_reason = None
    if not end:
        for candidate_name in reversed(INBOUND_STAGES[INBOUND_STAGE_ORDER[start_name] + 1:INBOUND_STAGE_ORDER[end_name]]):
            candidate_stamp = stamps.get(candidate_name)
            if candidate_stamp:
                end = candidate_stamp
                end_name_used = candidate_name
                fallback_reason = f"{end_name} ausente; usada a etapa anterior {candidate_name}"
                break
    base = {"ruleId": rule["ruleId"], "stageFrom": start_name, "stageTo": end_name, "startAt": start.isoformat(timespec="seconds"), "dueAt": due.isoformat(timespec="seconds")}
    if end:
        if end < start:
            return {**base, "stageToUsed": end_name_used, "fallbackApplied": bool(fallback_reason), "fallbackReason": fallback_reason, "status": "Not Calculated", "performance": None, "reason": "Evento final anterior ao evento inicial", "endAt": end.isoformat(timespec="seconds")}
        on_time = end <= due
        return {**base, "stageToUsed": end_name_used, "fallbackApplied": bool(fallback_reason), "fallbackReason": fallback_reason, "status": "Completed On Time" if on_time else "Completed Delay", "performance": "On Time" if on_time else "Delay", "reason": None, "endAt": end.isoformat(timespec="seconds")}
    later_exists = any(stamps.get(name) for name in INBOUND_STAGES[INBOUND_STAGE_ORDER[end_name] + 1:])
    if later_exists:
        return {**base, "status": "Not Calculated", "performance": None, "reason": f"Evento final ausente apesar de etapa posterior: {end_name}"}
    overdue = as_of > due
    return {**base, "status": "Open Delay" if overdue else "Open On Track", "performance": "Delay" if overdue else None, "reason": None}


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"Arquivo fonte não encontrado: {SOURCE}")
    aliases, sla_targets = load_sla_configuration()
    operational = OperationalConfiguration()
    calendar = BusinessCalendar(operational)
    if not operational.outbound_sources:
        raise ValueError("nenhuma fonte ativa cadastrada em FontesOutbound")
    if not operational.inbound_sources:
        raise ValueError("nenhuma fonte ativa cadastrada em FontesInbound")

    aggregates: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    daily_aggregates: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    stage_buckets: dict[tuple[str, str, str, str, str, str], dict[str, Any]] = {}
    daily_stage_buckets: dict[tuple[str, str, str, str, str, str], dict[str, Any]] = {}
    process_buckets: dict[tuple[str, str, str, str, str], dict[str, Any]] = {}
    daily_process_buckets: dict[tuple[str, str, str, str, str], dict[str, Any]] = {}
    combined_process_buckets: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    daily_combined_process_buckets: dict[tuple[str, str, str, str], dict[str, Any]] = {}
    clients: dict[str, dict[str, Any]] = {}
    details: dict[str, list[dict[str, Any]]] = defaultdict(list)
    capacity_groups: dict[tuple[str, str, str, date], list[dict[str, Any]]] = defaultdict(list)
    daily_capacity: list[dict[str, Any]] = []
    quality: dict[str, Any] = {"sourceRows": {"Outbound": 0, "OutboundLegacy": 0, "OutboundMigrated": 0, "Inbound": 0, "InboundLegacy": 0, "InboundMigrated": 0}, "cancelledOutbound": 0, "invalidCreationDates": {"Outbound": 0, "Inbound": 0}, "missingLines": {"Outbound": 0, "Inbound": 0}, "unclassified": {"Outbound": 0, "Inbound": 0}, "excludedInboundBMW": 0, "timestampWithoutDate": defaultdict(int), "operationalConfigFound": OPERATIONAL_CONFIG.exists(), "outboundFiles": [], "inboundFiles": [], "newSourceCancellationField": False}

    source_files = sorted({OUTBOUND_CSV_DIR / cfg["filename"] for cfg in operational.outbound_sources.values()})
    missing = [path.name for path in source_files if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"fontes Outbound ausentes: {', '.join(missing)}")
    inbound_source_files = sorted({INBOUND_CSV_DIR / cfg["filename"] for cfg in operational.inbound_sources.values()})
    inbound_missing = [path.name for path in inbound_source_files if not path.is_file()]
    if inbound_missing:
        raise FileNotFoundError(f"fontes Inbound ausentes: {', '.join(inbound_missing)}")
    source_as_of = max([SOURCE.stat().st_mtime, *[path.stat().st_mtime for path in source_files], *[path.stat().st_mtime for path in inbound_source_files]])
    source_as_of_dt = datetime.fromtimestamp(source_as_of)

    def ensure_client(key: str, name: str) -> dict[str, Any]:
        return clients.setdefault(key, {"key": key, "name": name, "operations": set(), "owners": {}})

    def process_outbound(*, key: str, name: str, owner_key: str, owner_name: str, owner_raw: str, created_day: date, lines: int | None, perf: str, stamps: dict[str, datetime | None], process_id: str, technical_key: str, type_status: str, source_kind: str) -> None:
        client = ensure_client(key, name)
        client["operations"].add("Outbound")
        client["owners"][owner_key] = {"key": owner_key, "name": owner_name}
        month = created_day.strftime("%Y-%m")
        metric_key = (key, "Outbound", month, owner_key)
        bucket = aggregates.setdefault(metric_key, new_metric(key, "Outbound", month, owner_key, source_kind))
        before_counts = {field: value for field, value in bucket.items() if isinstance(value, (int, float))}
        bucket["records"] += 1
        if source_kind == "csv":
            bucket["newSourceRecords"] += 1
        else:
            bucket["openMetricsAvailableRecords"] += 1
        if lines is None or lines <= 0:
            bucket["missingLinesRecords"] += 1
            quality["missingLines"]["Outbound"] += 1
        if perf == "Unclassified":
            bucket["unclassifiedRecords"] += 1
            quality["unclassified"]["Outbound"] += 1
        valid_lines = lines is not None and lines > 0
        if perf in {"On Time", "Delay"} and valid_lines:
            bucket["eligibleRecords"] += 1
            bucket["validLines"] += lines
            bucket["onTimeLines" if perf == "On Time" else "delayedLines"] += lines

        gross = classify_gross(operational, calendar, key, stamps, created_day, source_as_of_dt, owner_key)
        gross_status, gross_perf = gross["status"], gross["performance"]
        if gross.get("fallbackApplied"):
            bucket["calcFallbackRecords"] += 1
            if valid_lines:
                bucket["calcFallbackLines"] += lines
        if gross_status.startswith("Completed"):
            bucket["calcCompletedRecords"] += 1
            bucket["calcGrossEligibleRecords"] += 1
            if valid_lines:
                field = "OnTime" if gross_perf == "On Time" else "Delayed"
                bucket["calcCompletedEligibleLines"] += lines
                bucket["calcGrossEligibleLines"] += lines
                bucket[f"calcCompleted{field}Lines"] += lines
                bucket[f"calcGross{field}Lines"] += lines
        elif gross_status == "Open Delay":
            bucket["calcOpenOverdueRecords"] += 1
            bucket["calcGrossEligibleRecords"] += 1
            if valid_lines:
                bucket["calcOpenOverdueLines"] += lines
                bucket["calcGrossEligibleLines"] += lines
                bucket["calcGrossDelayedLines"] += lines
        elif gross_status == "Open On Track":
            bucket["calcOpenWithinRecords"] += 1
            if valid_lines:
                bucket["calcOpenWithinLines"] += lines
        else:
            bucket["calcNotCalculatedRecords"] += 1
            if valid_lines:
                bucket["calcNotCalculatedLines"] += lines
        if valid_lines and perf in {"On Time", "Delay"} and gross_perf in {"On Time", "Delay"}:
            bucket["comparisonEligibleLines"] += lines
            if perf == gross_perf == "On Time":
                bucket["comparisonBothOnTimeLines"] += lines
            elif perf == gross_perf == "Delay":
                bucket["comparisonBothDelayLines"] += lines
            elif perf == "On Time":
                bucket["comparisonSourceOnTimeCalcDelayLines"] += lines
            else:
                bucket["comparisonSourceDelayCalcOnTimeLines"] += lines

        for start_name, end_name in TRANSITIONS:
            period_buckets = (
                stage_buckets.setdefault(("Outbound", key, owner_key, month, start_name, end_name), new_stage("Outbound", key, month, start_name, end_name, owner_key)),
                daily_stage_buckets.setdefault(("Outbound", key, owner_key, created_day.isoformat(), start_name, end_name), new_stage("Outbound", key, created_day.isoformat(), start_name, end_name, owner_key)),
            )
            for stage in period_buckets:
                stage["totalRecords"] += 1
                start, end = stamps[start_name], stamps[end_name]
                if source_kind == "csv":
                    stage["performanceSourceRecords"] += 1
                    stage_result = classify_stage_performance(
                        calendar, key, stamps, start_name, end_name,
                        operational.stage_target(key, start_name, end_name, created_day, owner_key),
                        source_as_of_dt, owner_key,
                    )
                    if valid_lines:
                        stage["performanceTotalRecords"] += 1
                        stage["performanceTotalLines"] += lines
                        if stage_result["status"] in {"Completed On Time", "Completed Delay", "Open Delay"}:
                            stage["performanceEligibleRecords"] += 1
                            stage["performanceEligibleLines"] += lines
                            if stage_result["performance"] == "On Time":
                                stage["performanceOnTimeLines"] += lines
                            elif stage_result["performance"] == "Delay":
                                stage["performanceDelayedLines"] += lines
                        elif stage_result["status"] == "Open On Track":
                            stage["performanceOpenWithinLines"] += lines
                        else:
                            stage["performanceNotCalculatedLines"] += lines
                    if stage_result["status"] == "No SLA":
                        stage["performanceNoSlaRecords"] += 1
                if start and end:
                    if end < start:
                        stage["invalidSequenceCount"] += 1
                    else:
                        calendar_hours = (end - start).total_seconds() / 3600
                        business_hours = calendar.business_hours(key, start, end, owner_key)
                        stage["validCount"] += 1
                        stage["calendarValues"].append(calendar_hours)
                        stage["businessValues"].append(business_hours)
                        if calendar_hours == 0:
                            stage["zeroCount"] += 1
                elif start and not end:
                    end_position = next(i for i, item in enumerate(STAGES) if item[0] == end_name)
                    if any(stamps[stage_name] for stage_name, _, _ in STAGES[end_position + 1:]):
                        stage["incompleteCount"] += 1
                    else:
                        stage["openCount"] += 1
                        if source_kind != "csv" and source_as_of_dt >= start:
                            stage["openAgeValues"].append(calendar.business_hours(key, start, source_as_of_dt, owner_key))
                elif end and not start:
                    stage["incompleteCount"] += 1
                else:
                    stage["missingCount"] += 1

        # Process cards use explicit start/end transition boundaries from the
        # ProcessosOutbound sheet and only the consolidated source.
        if source_kind == "csv":
            process_results: dict[str, dict[str, Any]] = {}
            for rule in operational.process_rules_for(key, created_day, owner_key):
                process_buckets_for_period = (
                    process_buckets.setdefault(("Outbound", key, owner_key, month, rule["processKey"]), new_process("Outbound", key, month, rule, owner_key)),
                    daily_process_buckets.setdefault(("Outbound", key, owner_key, created_day.isoformat(), rule["processKey"]), new_process("Outbound", key, created_day.isoformat(), rule, owner_key)),
                )
                result = classify_process_performance(calendar, key, stamps, rule, source_as_of_dt, owner_key, operational, created_day)
                process_results[rule["processKey"]] = result
                for process in process_buckets_for_period:
                    process["deadlineSource"] = result.get("deadlineSource", "MetaHorasUteis")
                    process["deadlineType"] = result.get("deadlineType")
                    if result.get("deadlineSource") == "Gross":
                        process["targetHours"] = None
                    process["totalRecords"] += 1
                    if not valid_lines:
                        continue
                    process["totalLines"] += lines
                    if result["status"] in {"Completed On Time", "Completed Delay", "Open Delay"}:
                        process["eligibleRecords"] += 1
                        process["eligibleLines"] += lines
                        process["onTimeLines" if result["performance"] == "On Time" else "delayedLines"] += lines
                    elif result["status"] == "Open On Track":
                        process["openWithinLines"] += lines
                    else:
                        process["notCalculatedLines"] += lines
                    if result["status"] == "No SLA":
                        process["noSlaRecords"] += 1

            # A combined process is only eligible when both configured steps
            # are eligible for the same order/line. This preserves the
            # line-level intersection instead of adding independent totals.
            if {"PROCESSAMENTO", "SEPARACAO"}.issubset(process_results):
                combined_periods = (
                    combined_process_buckets.setdefault(("Outbound", key, owner_key, month), new_combined_process("Outbound", key, month, owner_key)),
                    daily_combined_process_buckets.setdefault(("Outbound", key, owner_key, created_day.isoformat()), new_combined_process("Outbound", key, created_day.isoformat(), owner_key)),
                )
                combined_result = process_results["PROCESSAMENTO"], process_results["SEPARACAO"]
                for combined in combined_periods:
                    combined["totalRecords"] += 1
                    if not valid_lines:
                        continue
                    combined["totalLines"] += lines
                    statuses = {item["status"] for item in combined_result}
                    if statuses.issubset({"Completed On Time", "Completed Delay", "Open Delay"}):
                        combined["eligibleRecords"] += 1
                        combined["eligibleLines"] += lines
                        if all(item["performance"] == "On Time" for item in combined_result):
                            combined["onTimeLines"] += lines
                        else:
                            combined["delayedLines"] += lines
                    elif statuses == {"Open On Track"}:
                        combined["openWithinLines"] += lines
                    else:
                        combined["notCalculatedLines"] += lines
                    if "No SLA" in statuses:
                        combined["noSlaRecords"] += 1

        detail_item = None
        if perf != "On Time" or gross_status != "Completed On Time" or gross.get("fallbackApplied"):
            detail_item = {"operation": "Outbound", "ownerKey": owner_key, "ownerName": owner_name, "ownerRaw": owner_raw, "sourceKind": source_kind, "sourcePerformanceAvailable": source_kind != "csv", "processId": process_id or "Sem identificador", "technicalKey": technical_key, "creationDate": created_day.isoformat(), "lineCount": lines, "typeStatus": type_status or "Não informado", "performance": perf, "grossStatus": gross_status, "grossReason": gross.get("reason"), "grossFallbackReason": gross.get("fallbackReason"), "grossRuleId": gross.get("ruleId"), "grossStageFrom": gross.get("stageFrom"), "grossStageTo": gross.get("stageTo"), "grossStageToUsed": gross.get("stageToUsed"), "grossStartAt": gross.get("startAt"), "grossDueAt": gross.get("dueAt"), "grossEndAt": gross.get("endAt"), "netStatus": "Não aplicável" if gross_perf not in {"On Time", "Delay"} else "Aguardando capacidade"}
            details[key].append(detail_item)

        capacity_stamp = stamps.get("Allocated") or stamps.get("Released")
        if capacity_stamp and valid_lines:
            normalized = calendar.normalize_release(key, capacity_stamp, owner_key)
            capacity_groups[("Outbound", key, owner_key, normalized.date())].append({"metricKey": metric_key, "dailyMetricKey": (key, "Outbound", created_day.isoformat(), owner_key), "month": month, "lines": lines, "performance": gross_perf, "released": capacity_stamp, "created": stamps["Creation"] or datetime.combine(created_day, time.min), "technicalKey": technical_key, "detail": detail_item})
        elif gross_perf in {"On Time", "Delay"} and valid_lines:
            bucket["netUnavailableLines"] += lines
            if detail_item:
                detail_item["netStatus"] = "NET não calculado: alocação/liberação ausente"

        daily_key = (key, "Outbound", created_day.isoformat(), owner_key)
        daily_bucket = daily_aggregates.setdefault(daily_key, new_metric(key, "Outbound", created_day.isoformat(), owner_key, source_kind))
        for field, before_value in before_counts.items():
            daily_bucket[field] += bucket.get(field, 0) - before_value

    def process_inbound(*, key: str, name: str, owner_key: str, owner_name: str, owner_raw: str, created_day: date, lines: int | None, source_perf: str, stamps: dict[str, datetime | None] | None, source_due: datetime | None, process_id: str, technical_key: str, type_status: str, source_kind: str) -> None:
        client = ensure_client(key, name)
        client["operations"].add("Inbound")
        client["owners"][owner_key] = {"key": owner_key, "name": owner_name}
        month = created_day.strftime("%Y-%m")
        metric_key = (key, "Inbound", month, owner_key)
        bucket = aggregates.setdefault(metric_key, new_metric(key, "Inbound", month, owner_key, source_kind))
        before_counts = {field: value for field, value in bucket.items() if isinstance(value, (int, float))}
        bucket["records"] += 1
        if source_kind == "csv":
            bucket["newSourceRecords"] += 1
            result = classify_inbound(operational, calendar, key, stamps or {}, source_due, created_day, source_as_of_dt, owner_key)
            perf = result["performance"] or "Unclassified"
        else:
            legacy_status = "Completed On Time" if source_perf == "On Time" else "Completed Delay" if source_perf == "Delay" else "Not Calculated"
            result = {"status": legacy_status, "performance": source_perf if source_perf in {"On Time", "Delay"} else None, "reason": None, "stageFrom": "Creation", "stageTo": "Finish"}
            perf = source_perf
        valid_lines = lines is not None and lines > 0
        if not valid_lines:
            bucket["missingLinesRecords"] += 1
            quality["missingLines"]["Inbound"] += 1
        if perf == "Unclassified":
            bucket["unclassifiedRecords"] += 1
            quality["unclassified"]["Inbound"] += 1
        if perf in {"On Time", "Delay"} and valid_lines:
            bucket["eligibleRecords"] += 1
            bucket["validLines"] += lines
            bucket["onTimeLines" if perf == "On Time" else "delayedLines"] += lines
        gross_perf = result.get("performance")
        if result.get("fallbackApplied"):
            bucket["calcFallbackRecords"] += 1
            if valid_lines:
                bucket["calcFallbackLines"] += lines
        if result["status"].startswith("Completed"):
            bucket["calcCompletedRecords"] += 1
            bucket["calcGrossEligibleRecords"] += 1
            if valid_lines:
                field = "OnTime" if gross_perf == "On Time" else "Delayed"
                bucket["calcCompletedEligibleLines"] += lines
                bucket["calcGrossEligibleLines"] += lines
                bucket[f"calcCompleted{field}Lines"] += lines
                bucket[f"calcGross{field}Lines"] += lines
        elif result["status"] == "Open Delay":
            bucket["calcOpenOverdueRecords"] += 1
            bucket["calcGrossEligibleRecords"] += 1
            if valid_lines:
                bucket["calcOpenOverdueLines"] += lines
                bucket["calcGrossEligibleLines"] += lines
                bucket["calcGrossDelayedLines"] += lines
        elif result["status"] == "Open On Track":
            bucket["calcOpenWithinRecords"] += 1
            if valid_lines:
                bucket["calcOpenWithinLines"] += lines
        else:
            bucket["calcNotCalculatedRecords"] += 1
            if valid_lines:
                bucket["calcNotCalculatedLines"] += lines

        inbound_stamps = stamps or {}
        for start_name, end_name in INBOUND_TRANSITIONS:
            period_buckets = (
                stage_buckets.setdefault(("Inbound", key, owner_key, month, start_name, end_name), new_stage("Inbound", key, month, start_name, end_name, owner_key)),
                daily_stage_buckets.setdefault(("Inbound", key, owner_key, created_day.isoformat(), start_name, end_name), new_stage("Inbound", key, created_day.isoformat(), start_name, end_name, owner_key)),
            )
            for stage in period_buckets:
                stage["totalRecords"] += 1
                start, end = inbound_stamps.get(start_name), inbound_stamps.get(end_name)
                if start and end:
                    if end < start:
                        stage["invalidSequenceCount"] += 1
                    else:
                        calendar_hours = (end - start).total_seconds() / 3600
                        stage["validCount"] += 1
                        stage["calendarValues"].append(calendar_hours)
                        stage["businessValues"].append(calendar.business_hours(key, start, end, owner_key))
                        if calendar_hours == 0:
                            stage["zeroCount"] += 1
                elif start and not end:
                    end_position = INBOUND_STAGE_ORDER[end_name]
                    if any(inbound_stamps.get(stage_name) for stage_name in INBOUND_STAGES[end_position + 1:]):
                        stage["incompleteCount"] += 1
                    else:
                        stage["openCount"] += 1
                        if source_as_of_dt >= start:
                            stage["openAgeValues"].append(calendar.business_hours(key, start, source_as_of_dt, owner_key))
                elif end and not start:
                    stage["incompleteCount"] += 1
                else:
                    stage["missingCount"] += 1

        detail_item = None
        if perf != "On Time" or result.get("fallbackApplied"):
            detail_item = {
                "operation": "Inbound", "ownerKey": owner_key, "ownerName": owner_name, "ownerRaw": owner_raw,
                "sourceKind": source_kind, "sourcePerformanceAvailable": source_kind != "csv", "processId": process_id or "Sem identificador",
                "technicalKey": technical_key, "creationDate": created_day.isoformat(), "lineCount": lines,
                "typeStatus": type_status or "Não informado", "performance": perf, "grossStatus": result.get("status"),
                "grossReason": result.get("reason"), "grossFallbackReason": result.get("fallbackReason"),
                "grossRuleId": result.get("ruleId"), "grossStageFrom": result.get("stageFrom"), "grossStageTo": result.get("stageTo"),
                "grossStageToUsed": result.get("stageToUsed"), "grossStartAt": result.get("startAt"),
                "grossDueAt": result.get("dueAt"), "grossEndAt": result.get("endAt"), "netStatus": "Não aplicável" if gross_perf not in {"On Time", "Delay"} else "Aguardando capacidade",
            }
            details[key].append(detail_item)

        capacity_stage = result.get("stageFrom") or "Creation"
        capacity_stamp = inbound_stamps.get(capacity_stage) or inbound_stamps.get("Creation")
        if capacity_stamp and valid_lines:
            normalized = calendar.normalize_release(key, capacity_stamp, owner_key)
            capacity_groups[("Inbound", key, owner_key, normalized.date())].append({"metricKey": metric_key, "dailyMetricKey": (key, "Inbound", created_day.isoformat(), owner_key), "month": month, "lines": lines, "performance": gross_perf, "released": capacity_stamp, "created": inbound_stamps.get("Creation") or datetime.combine(created_day, time.min), "technicalKey": technical_key, "detail": detail_item})
        elif gross_perf in {"On Time", "Delay"} and valid_lines:
            bucket["netUnavailableLines"] += lines
            if detail_item:
                detail_item["netStatus"] = "NET não calculado: evento inicial ausente"
        daily_key = (key, "Inbound", created_day.isoformat(), owner_key)
        daily_bucket = daily_aggregates.setdefault(daily_key, new_metric(key, "Inbound", created_day.isoformat(), owner_key, source_kind))
        for field, before_value in before_counts.items():
            daily_bucket[field] += bucket.get(field, 0) - before_value

    book = load_workbook(SOURCE, read_only=True, data_only=True)
    rows = book["OverviewOrder"].iter_rows(values_only=True)
    headers = {clean(name): i for i, name in enumerate(next(rows))}
    stage_indexes = [(name, headers[day], headers[clock]) for name, day, clock in STAGES]
    for row in rows:
        quality["sourceRows"]["Outbound"] += 1
        quality["sourceRows"]["OutboundLegacy"] += 1
        if is_cancelled(row[headers["CANCELLED"]]):
            quality["cancelledOutbound"] += 1
            continue
        person = identity(row[headers["Client ID"]], aliases)
        key = person["key"]
        created_day = as_date(row[headers["CREATION_DATE"]])
        if not created_day:
            quality["invalidCreationDates"]["Outbound"] += 1
            continue
        source_config = operational.outbound_sources.get(key)
        if source_config and created_day >= source_config["cutoff"]:
            continue
        stamps: dict[str, datetime | None] = {}
        for stage_name, date_index, time_index in stage_indexes:
            date_value, time_value = row[date_index], row[time_index]
            if as_time(time_value) and not as_date(date_value):
                quality["timestampWithoutDate"][stage_name] += 1
            stamps[stage_name] = combined(date_value, time_value)
        process_outbound(key=key, name=person["name"], owner_key="HISTORICO", owner_name="Histórico", owner_raw="", created_day=created_day, lines=number(row[headers["NUM_LINES"]]), perf=performance(row[headers["Performance"]]), stamps=stamps, process_id=clean(row[headers["ORDER_ID"]]), technical_key=clean(row[headers.get("Chaveamento", headers["ORDER_ID"])]), type_status=clean(row[headers["ORDER_TYPE2"]]), source_kind="legacy")

    seen_keys: set[str] = set()
    for path in source_files:
        file_configs = {key: config for key, config in operational.outbound_sources.items() if config["filename"].lower() == path.name.lower()}
        audits = {key: {"clientKey": key, "file": path.name, "modifiedAt": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"), "rows": 0, "validRows": 0, "cancelledRows": 0, "duplicates": 0, "minCreation": None, "maxCreation": None, "status": "Válido", "validation": "Base consolidada, schema, chave, Owner, cobertura e timestamps válidos"} for key in file_configs}
        creation_dates: list[date] = []
        dates_by_client: dict[str, list[date]] = defaultdict(list)
        covered_clients: set[str] = set()
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle, delimiter="|")
            columns = set(reader.fieldnames or [])
            missing_columns = sorted(CSV_REQUIRED_COLUMNS - columns)
            if missing_columns:
                raise ValueError(f"{path.name}: colunas obrigatórias ausentes: {', '.join(missing_columns)}")
            quality["newSourceCancellationField"] = quality["newSourceCancellationField"] or "CANCELLED" in columns
            for line_number, raw in enumerate(reader, 2):
                person = identity(raw.get("Client"), aliases)
                key = person["key"]
                config = file_configs.get(key)
                if not config:
                    raise ValueError(f"{path.name}:{line_number}: cliente {key} não está ativo em FontesOutbound")
                audit = audits[key]
                audit["rows"] += 1
                if "CANCELLED" in columns and is_cancelled(raw.get("CANCELLED")):
                    audit["cancelledRows"] += 1
                    quality["cancelledOutbound"] += 1
                    continue
                if config["sourceAlias"] and clean(raw.get("Client")).upper() != config["sourceAlias"]:
                    raise ValueError(f"{path.name}:{line_number}: alias {clean(raw.get('Client'))} difere de {config['sourceAlias']} para {key}")
                raw_owner = clean(raw.get("Owner"))
                if config["ownerRequired"] and not raw_owner:
                    raise ValueError(f"{path.name}:{line_number}: Owner vazio")
                try:
                    stamps = {stage: csv_timestamp(raw.get(column)) for stage, column in CSV_STAGE_COLUMNS.items()}
                except ValueError as exc:
                    raise ValueError(f"{path.name}:{line_number}: {exc}") from exc
                created = stamps["Creation"]
                if not created:
                    raise ValueError(f"{path.name}:{line_number}: Creation Date vazio")
                technical_key = f"{clean(raw.get('Site'))}|{clean(raw.get('Client'))}|{clean(raw.get('Order Id'))}"
                if technical_key in seen_keys:
                    audit["duplicates"] += 1
                    raise ValueError(f"{path.name}:{line_number}: chave duplicada {technical_key}")
                seen_keys.add(technical_key)
                owner = operational.owner_identity(key, raw_owner)
                if owner["key"] == "NÃO INFORMADO":
                    raise ValueError(f"{path.name}:{line_number}: Owner não informado")
                creation_dates.append(created.date())
                dates_by_client[key].append(created.date())
                covered_clients.add(key)
                audit["validRows"] += 1
                quality["sourceRows"]["Outbound"] += 1
                quality["sourceRows"]["OutboundMigrated"] += 1
                if created.date() < config["cutoff"]:
                    continue
                process_outbound(key=key, name=person["name"], owner_key=owner["key"], owner_name=owner["name"], owner_raw=owner["raw"], created_day=created.date(), lines=number(raw.get("Lines")), perf="Unclassified", stamps=stamps, process_id=clean(raw.get("Order Id")), technical_key=technical_key, type_status=clean(raw.get("Order Type")) or clean(raw.get("Status")), source_kind="csv")
        if not creation_dates:
            raise ValueError(f"{path.name}: arquivo vazio")
        missing_clients = sorted(set(file_configs) - covered_clients)
        if missing_clients:
            raise ValueError(f"{path.name}: clientes configurados sem registros: {', '.join(missing_clients)}")
        for key, config in file_configs.items():
            client_dates = dates_by_client[key]
            if min(client_dates) > config["cutoff"]:
                raise ValueError(f"{path.name}: corte {config['cutoff'].isoformat()} sem cobertura para {key}; primeiro registro {min(client_dates).isoformat()}")
            audit = audits[key]
            audit["minCreation"] = min(client_dates).isoformat()
            audit["maxCreation"] = max(client_dates).isoformat()
            quality["outboundFiles"].append(audit)

    rows = book["OverviewAdvice"].iter_rows(values_only=True)
    headers = {clean(name): i for i, name in enumerate(next(rows))}
    for row in rows:
        quality["sourceRows"]["Inbound"] += 1
        quality["sourceRows"]["InboundLegacy"] += 1
        person = identity(row[headers["Client"]], aliases)
        key = person["key"]
        if key == "BMW":
            quality["excludedInboundBMW"] += 1
            continue
        created = as_date(row[headers["Creation"]])
        if not created:
            quality["invalidCreationDates"]["Inbound"] += 1
            continue
        source_config = operational.inbound_sources.get(key)
        if source_config and created >= source_config["cutoff"]:
            continue
        legacy_stamps = {
            "Creation": combined(row[headers["Creation"]], row[headers["Creation2"]]),
            "Released": combined(row[headers["Released"]], row[headers["Released2"]]),
            "In Progress": combined(row[headers["Start"]], row[headers["Start2"]]),
            "Finish": combined(row[headers["Finish"]], row[headers["Finish2"]]),
        }
        process_inbound(key=key, name=person["name"], owner_key="HISTORICO", owner_name="Histórico", owner_raw="", created_day=created, lines=number(row[headers["Lines"]]), source_perf=performance(row[headers["Performance"]]), stamps=legacy_stamps, source_due=None, process_id=clean(row[headers["Pre-Advice ID"]]), technical_key=clean(row[headers["Pre-Advice ID"]]), type_status=clean(row[headers["Status"]]), source_kind="legacy")

    inbound_seen_keys: set[str] = set()
    for path in inbound_source_files:
        file_configs = {key: config for key, config in operational.inbound_sources.items() if config["filename"].lower() == path.name.lower()}
        audit = {"file": path.name, "modifiedAt": datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds"), "rows": 0, "validRows": 0, "duplicates": 0, "minCreation": None, "maxCreation": None, "status": "Válido", "validation": "Schema, chave, Owner, cobertura e timestamps válidos"}
        creation_dates: list[date] = []
        dates_by_client: dict[str, list[date]] = defaultdict(list)
        covered_clients: set[str] = set()
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle, delimiter="|")
            columns = set(reader.fieldnames or [])
            missing_columns = sorted(INBOUND_REQUIRED_COLUMNS - columns)
            if missing_columns:
                raise ValueError(f"{path.name}: colunas obrigatórias ausentes: {', '.join(missing_columns)}")
            for line_number, raw in enumerate(reader, 2):
                audit["rows"] += 1
                person = identity(raw.get("Client"), aliases)
                key = person["key"]
                config = file_configs.get(key)
                if not config:
                    raise ValueError(f"{path.name}:{line_number}: cliente {key} não está ativo em FontesInbound")
                if config["sourceAlias"] and clean(raw.get("Client")).upper() != config["sourceAlias"]:
                    raise ValueError(f"{path.name}:{line_number}: alias {clean(raw.get('Client'))} difere de {config['sourceAlias']} para {key}")
                raw_owner = clean(raw.get("Owner"))
                if config["ownerRequired"] and not raw_owner:
                    raise ValueError(f"{path.name}:{line_number}: Owner vazio")
                try:
                    stamps = {stage: csv_timestamp(raw.get(column)) for stage, column in INBOUND_CSV_STAGE_COLUMNS.items()}
                    source_due = csv_timestamp(raw.get("Due Date"))
                except ValueError as exc:
                    raise ValueError(f"{path.name}:{line_number}: {exc}") from exc
                created = stamps["Creation"]
                if not created:
                    raise ValueError(f"{path.name}:{line_number}: Creation Date vazio")
                technical_key = f"{clean(raw.get('Site'))}|{clean(raw.get('Client'))}|{clean(raw.get('Pre-Advice ID'))}"
                if technical_key in inbound_seen_keys:
                    audit["duplicates"] += 1
                    raise ValueError(f"{path.name}:{line_number}: chave duplicada {technical_key}")
                inbound_seen_keys.add(technical_key)
                owner = operational.owner_identity(key, raw_owner)
                if owner["key"] == "NÃO INFORMADO":
                    raise ValueError(f"{path.name}:{line_number}: Owner não informado")
                creation_dates.append(created.date())
                dates_by_client[key].append(created.date())
                covered_clients.add(key)
                audit["validRows"] += 1
                quality["sourceRows"]["Inbound"] += 1
                quality["sourceRows"]["InboundMigrated"] += 1
                if created.date() < config["cutoff"]:
                    continue
                process_inbound(key=key, name=person["name"], owner_key=owner["key"], owner_name=owner["name"], owner_raw=owner["raw"], created_day=created.date(), lines=number(raw.get("Lines")), source_perf="Unclassified", stamps=stamps, source_due=source_due, process_id=clean(raw.get("Pre-Advice ID")), technical_key=technical_key, type_status=clean(raw.get("Type")) or clean(raw.get("Status")), source_kind="csv")
        if not audit["rows"]:
            raise ValueError(f"{path.name}: arquivo vazio")
        if not creation_dates:
            raise ValueError(f"{path.name}: nenhuma Creation Date válida")
        missing_clients = sorted(set(file_configs) - covered_clients)
        if missing_clients:
            raise ValueError(f"{path.name}: clientes configurados sem registros: {', '.join(missing_clients)}")
        for key, config in file_configs.items():
            client_dates = dates_by_client[key]
            if min(client_dates) > config["cutoff"]:
                raise ValueError(f"{path.name}: corte {config['cutoff'].isoformat()} sem cobertura para {key}")
        audit["clients"] = sorted(covered_clients)
        audit["minCreation"] = min(creation_dates).isoformat()
        audit["maxCreation"] = max(creation_dates).isoformat()
        quality["inboundFiles"].append(audit)
    book.close()
    inventory = load_inventory(aliases, operational)
    for inventory_row in inventory["latestRows"]:
        if inventory_row["clientKey"] == "SEM_CLIENTE":
            continue
        client = ensure_client(inventory_row["clientKey"], inventory_row["clientName"])
        client["operations"].add("Inventory")
        client["owners"][inventory_row["ownerKey"]] = {"key": inventory_row["ownerKey"], "name": inventory_row["ownerName"]}

    for (operation, key, owner, business_day), records in sorted(capacity_groups.items()):
        records.sort(key=lambda item: (item["released"], item["created"], item["technicalKey"]))
        official = operational.inbound_capacity(key, business_day, owner) if operation == "Inbound" else operational.capacity(key, business_day, owner)
        profile = operational.inbound_capacity_profile(key, owner) if operation == "Inbound" else operational.capacity_profile(key, owner)
        suggested = profile.get("suggestedCapacity")
        remaining = official
        total_lines = sum(item["lines"] for item in records)
        included_total = excluded_total = 0
        allocations = allocate_inbound_net_capacity(records, official) if operation == "Inbound" else None
        for index, record in enumerate(records):
            bucket = aggregates[record["metricKey"]]
            daily_bucket = daily_aggregates[record["dailyMetricKey"]]
            if remaining is None:
                if record["performance"] in {"On Time", "Delay"}:
                    bucket["netUnavailableLines"] += record["lines"]
                    daily_bucket["netUnavailableLines"] += record["lines"]
                if record["detail"] and record["performance"] == "Delay":
                    record["detail"]["netStatus"] = "NET não calculado: capacidade oficial ausente"
                continue
            if allocations is not None:
                included, excluded = allocations[index]
            else:
                included, excluded, remaining = split_capacity(record["lines"], remaining)
            included_total += included
            excluded_total += excluded
            if record["performance"] in {"On Time", "Delay"}:
                bucket["netEligibleLines"] += included
                bucket["capacityExcludedLines"] += excluded
                bucket["netOnTimeLines" if record["performance"] == "On Time" else "netDelayedLines"] += included
                daily_bucket["netEligibleLines"] += included
                daily_bucket["capacityExcludedLines"] += excluded
                daily_bucket["netOnTimeLines" if record["performance"] == "On Time" else "netDelayedLines"] += included
            if record["detail"] and record["performance"] == "Delay":
                record["detail"].update({"netIncludedLines": included, "capacityExcludedLines": excluded, "netStatus": "Integralmente no NET" if excluded == 0 else "Fora do NET por capacidade" if included == 0 else "Parcialmente no NET"})
        daily_capacity.append({"operation": operation, "clientKey": key, "ownerKey": owner, "date": business_day.isoformat(), "officialCapacity": official, "suggestedCapacity": suggested, "releasedLines": total_lines, "netIncludedLines": included_total if official is not None else None, "excludedLines": excluded_total if official is not None else None, "utilization": included_total / official if official else None})

    stage_metrics: list[dict[str, Any]] = []
    for stage in stage_buckets.values():
        business, calendar_values, open_ages = stage.pop("businessValues"), stage.pop("calendarValues"), stage.pop("openAgeValues")
        valid = stage["validCount"]
        coverage = valid / stage["totalRecords"] if stage["totalRecords"] else None
        zero_rate = stage["zeroCount"] / valid if valid else None
        month_day = datetime.strptime(stage["month"] + "-01", "%Y-%m-%d").date()
        target = operational.stage_target(stage["clientKey"], stage["stageFrom"], stage["stageTo"], month_day, stage["ownerKey"]) if stage["operation"] == "Outbound" else None
        trust = "Sem confiabilidade" if stage["clientKey"] == "BMW" and stage["stageFrom"] == "Released" and stage["stageTo"] == "Allocated" else ("Confiável" if coverage is not None and coverage >= operational.rules["CoberturaConfiavel"] else "Atenção" if coverage is not None and coverage >= operational.rules["CoberturaAtencao"] else "Baixa cobertura")
        stage.update({"coverage": coverage, "meanCalendarHours": round(statistics.fmean(calendar_values), 2) if calendar_values else None, "medianCalendarHours": round(statistics.median(calendar_values), 2) if calendar_values else None, "p90CalendarHours": round(quantile(calendar_values, .90), 2) if calendar_values else None, "meanBusinessHours": round(statistics.fmean(business), 2) if business else None, "medianBusinessHours": round(statistics.median(business), 2) if business else None, "p90BusinessHours": round(quantile(business, .90), 2) if business else None, "medianBusinessDays": round(statistics.median(business) / operational.rules["HorasDiaUtil"], 2) if business else None, "medianOpenBusinessHours": round(statistics.median(open_ages), 2) if open_ages else None, "maxOpenBusinessHours": round(max(open_ages), 2) if open_ages else None, "zeroRate": zero_rate, "zeroWarning": zero_rate is not None and zero_rate > operational.rules["AlertaDuracaoZero"], "targetBusinessHours": target, "trust": trust})
        stage_metrics.append(stage)

    daily_stage_metrics: list[dict[str, Any]] = []
    for stage in daily_stage_buckets.values():
        business, calendar_values, open_ages = stage.pop("businessValues"), stage.pop("calendarValues"), stage.pop("openAgeValues")
        valid = stage["validCount"]
        coverage = valid / stage["totalRecords"] if stage["totalRecords"] else None
        zero_rate = stage["zeroCount"] / valid if valid else None
        period_day = datetime.strptime(stage["month"], "%Y-%m-%d").date()
        target = operational.stage_target(stage["clientKey"], stage["stageFrom"], stage["stageTo"], period_day, stage["ownerKey"]) if stage["operation"] == "Outbound" else None
        trust = "Sem confiabilidade" if stage["clientKey"] == "BMW" and stage["stageFrom"] == "Released" and stage["stageTo"] == "Allocated" else ("Confiável" if coverage is not None and coverage >= operational.rules["CoberturaConfiavel"] else "Atenção" if coverage is not None and coverage >= operational.rules["CoberturaAtencao"] else "Baixa cobertura")
        stage.update({"coverage": coverage, "meanCalendarHours": round(statistics.fmean(calendar_values), 2) if calendar_values else None, "medianCalendarHours": round(statistics.median(calendar_values), 2) if calendar_values else None, "p90CalendarHours": round(quantile(calendar_values, .90), 2) if calendar_values else None, "meanBusinessHours": round(statistics.fmean(business), 2) if business else None, "medianBusinessHours": round(statistics.median(business), 2) if business else None, "p90BusinessHours": round(quantile(business, .90), 2) if business else None, "medianBusinessDays": round(statistics.median(business) / operational.rules["HorasDiaUtil"], 2) if business else None, "medianOpenBusinessHours": round(statistics.median(open_ages), 2) if open_ages else None, "maxOpenBusinessHours": round(max(open_ages), 2) if open_ages else None, "zeroRate": zero_rate, "zeroWarning": zero_rate is not None and zero_rate > operational.rules["AlertaDuracaoZero"], "targetBusinessHours": target, "trust": trust})
        daily_stage_metrics.append({field: stage.get(field) for field in ("operation", "clientKey", "ownerKey", "month", "stageFrom", "stageTo", "totalRecords", "validCount", "missingCount", "openCount", "incompleteCount", "invalidSequenceCount", "zeroCount", "meanBusinessHours", "medianBusinessHours", "p90BusinessHours", "medianOpenBusinessHours", "targetBusinessHours", "trust", "performanceSourceRecords", "performanceTotalRecords", "performanceTotalLines", "performanceEligibleRecords", "performanceEligibleLines", "performanceOnTimeLines", "performanceDelayedLines", "performanceOpenWithinLines", "performanceNotCalculatedLines", "performanceNoSlaRecords")})

    process_metrics = sorted(process_buckets.values(), key=lambda item: (item["operation"], item["clientKey"], item["month"], item["ownerKey"], item["processKey"]))
    daily_process_metrics = sorted(daily_process_buckets.values(), key=lambda item: (item["operation"], item["clientKey"], item["month"], item["ownerKey"], item["processKey"]))
    combined_process_metrics = sorted(combined_process_buckets.values(), key=lambda item: (item["operation"], item["clientKey"], item["month"], item["ownerKey"]))
    daily_combined_process_metrics = sorted(daily_combined_process_buckets.values(), key=lambda item: (item["operation"], item["clientKey"], item["month"], item["ownerKey"]))
    process_rules = []
    for (client_key, owner_key, _), entries in operational.process_rules.items():
        for entry in entries:
            process_rules.append({"clientKey": client_key, **{key: (value.isoformat() if isinstance(value, date) else value) for key, value in entry.items()}})

    for key, item_rows in details.items():
        selected: list[dict[str, Any]] = []
        by_owner: dict[str, list[dict[str, Any]]] = defaultdict(list)
        for item in item_rows:
            by_owner[item.get("ownerKey", "ALL")].append(item)
        for owner_rows in by_owner.values():
            owner_rows.sort(key=lambda item: (item["creationDate"], item.get("lineCount") or 0), reverse=True)
            selected.extend(owner_rows[:600])
        selected.sort(key=lambda item: (item["creationDate"], item.get("lineCount") or 0), reverse=True)
        details[key] = selected

    client_rows = []
    capacity_profiles: list[dict[str, Any]] = []
    for key, item in sorted(clients.items(), key=lambda pair: pair[1]["name"]):
        owners = sorted(item["owners"].values(), key=lambda value: (value["key"] == "HISTORICO", value["name"]))
        for owner in owners:
            capacity_profiles.append(operational.capacity_profile(key, owner["key"]))
            capacity_profiles.append(operational.inbound_capacity_profile(key, owner["key"]))
        client_rows.append({"key": key, "name": item["name"], "operations": sorted(item["operations"]), "owners": owners, "targets": {op: sla_targets.get((key, op), {"target": .95, "warning": .90}) for op in ("Outbound", "Inbound")}, "capacity": operational.capacity_profile(key, "DEFAULT"), "inboundCapacity": operational.inbound_capacity_profile(key, "DEFAULT"), "grossRule": operational.gross_rule_summary(key), "inboundRule": operational.inbound_rule_summary(key), "outboundSource": ({"cutoff": operational.outbound_sources[key]["cutoff"].isoformat(), "file": operational.outbound_sources[key]["filename"]} if key in operational.outbound_sources else None), "inboundSource": ({"cutoff": operational.inbound_sources[key]["cutoff"].isoformat(), "file": operational.inbound_sources[key]["filename"]} if key in operational.inbound_sources else None)})
    payload = {
        "meta": {"source": f"{SOURCE.name} + {OUTBOUND_CSV_DIR.relative_to(ROOT)} + {INBOUND_CSV_DIR.relative_to(ROOT)}", "sourceUpdatedAt": source_as_of_dt.isoformat(timespec="seconds"), "generatedAt": datetime.now().isoformat(timespec="seconds"), "defaultClient": "GWM" if "GWM" in clients else client_rows[0]["key"], "defaultMonthCount": 3, "operationalConfig": str(OPERATIONAL_CONFIG.relative_to(ROOT))},
        "clients": client_rows,
        "metrics": sorted(aggregates.values(), key=lambda item: (item["clientKey"], item["month"], item["operation"], item["ownerKey"])),
        "dailyMetrics": sorted(daily_aggregates.values(), key=lambda item: (item["clientKey"], item["month"], item["operation"], item["ownerKey"])),
        "stageMetrics": sorted(stage_metrics, key=lambda item: (item["operation"], item["clientKey"], item["month"], item["ownerKey"], (TRANSITIONS if item["operation"] == "Outbound" else INBOUND_TRANSITIONS).index((item["stageFrom"], item["stageTo"])))),
        "dailyStageMetrics": sorted(daily_stage_metrics, key=lambda item: (item["operation"], item["clientKey"], item["month"], item["ownerKey"], (TRANSITIONS if item["operation"] == "Outbound" else INBOUND_TRANSITIONS).index((item["stageFrom"], item["stageTo"])))),
        "processMetrics": process_metrics,
        "dailyProcessMetrics": daily_process_metrics,
        "combinedProcessMetrics": combined_process_metrics,
        "dailyCombinedProcessMetrics": daily_combined_process_metrics,
        "processRules": sorted(process_rules, key=lambda item: (item["clientKey"], item["ownerKey"], item["processKey"], item["from"])),
        "dailyCapacity": daily_capacity,
        "capacityProfiles": capacity_profiles,
        "details": details,
        "inventory": inventory,
        "quality": {**quality, "timestampWithoutDate": dict(quality["timestampWithoutDate"]), "rules": operational.rules},
    }
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    temp_output = OUTPUT.with_suffix(".json.tmp")
    temp_output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    os.replace(temp_output, OUTPUT)
    print(f"Gerado: {OUTPUT}")
    print(f"Clientes: {len(client_rows)} | Métricas: {len(payload['metrics'])} | Etapas: {len(stage_metrics)} | Tamanho: {OUTPUT.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
