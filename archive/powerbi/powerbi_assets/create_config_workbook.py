from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(r"C:\performance_geral_cjm\00 - HTML's update")
OUTPUT = ROOT / "MetasSLA.xlsx"

aliases = [
    ("BRALGAR", "ALGAR", "ALGAR", "Both"),
    ("BRASUS", "ASUS", "ASUS", "Outbound"),
    ("BRACBZ", "ACBZ", "ACBZ", "Inbound"),
    ("BRBAT", "BAT", "BAT", "Both"),
    ("BRBMW", "BMW", "BMW", "Both"),
    ("BRBT", "BT", "BT", "Both"),
    ("BRCORBION", "CORBION", "CORBION", "Both"),
    ("BRDUCATI", "DUCATI", "DUCATI", "Outbound"),
    ("BRDUCATIMO", "DUCATI", "DUCATI", "Inbound"),
    ("BREDELWHITE", "EDELWHITE", "EDELWHITE", "Outbound"),
    ("BREDELWHIT", "EDELWHITE", "EDELWHITE", "Inbound"),
    ("BRFIVEHANDS", "FIVEHANDS", "FIVEHANDS", "Outbound"),
    ("BRFIVEHNDS", "FIVEHANDS", "FIVEHANDS", "Inbound"),
    ("BRGAC", "GAC", "GAC", "Outbound"),
    ("BRGACMES", "GAC", "GAC", "Inbound"),
    ("BRGNUTRA", "GNUTRA", "GNUTRA", "Both"),
    ("BRGWM", "GWM", "GWM", "Both"),
    ("BRHANSEN", "HANSEN", "HANSEN", "Both"),
    ("BRHARLEY", "HARLEY", "HARLEY", "Outbound"),
    ("BRHARLD", "HARLEY", "HARLEY", "Inbound"),
    ("BRJETOUR", "JETOUR", "JETOUR", "Outbound"),
    ("BRCJ1JET", "JETOUR", "JETOUR", "Inbound"),
    ("BRMWC", "MWC", "MWC", "Both"),
    ("BROLESEN", "OLESEN", "OLESEN", "Both"),
    ("BROMODA", "OMODA", "OMODA", "Both"),
    ("BRROQUETTE", "ROQUETTE", "ROQUETTE", "Outbound"),
    ("BRRQTCAJ", "ROQUETTE", "ROQUETTE", "Inbound"),
    ("BRSAIC", "SAIC", "SAIC", "Both"),
    ("BRSCANDERRA", "SCANDERRA", "SCANDERRA", "Outbound"),
    ("BRSCANDERR", "SCANDERRA", "SCANDERRA", "Inbound"),
    ("BRCLARO", "CLARO", "CLARO", "Outbound"),
    ("BRCLAROSPO", "CLARO", "CLARO", "Inbound"),
    ("BRXSYS", "XSYS", "XSYS", "Outbound"),
    ("BRCJ1XYS", "XSYS", "XSYS", "Inbound"),
]

clients = sorted({(key, name) for _, key, name, _ in aliases})

wb = Workbook()
ws = wb.active
ws.title = "MetasSLA"
ws.append(["ClientKey", "Operation", "ValidFrom", "ValidTo", "TargetPct", "WarningPct"])
for key, _ in clients:
    ws.append([key, "Outbound", date(2025, 1, 1), None, 0.95, 0.90])
    if key != "BMW":
        ws.append([key, "Inbound", date(2025, 1, 1), None, 0.95, 0.90])

alias_ws = wb.create_sheet("ClientAliases")
alias_ws.append(["ClientAlias", "ClientKey", "ClientName", "Operation"])
for row in aliases:
    alias_ws.append(row)

for sheet, table_name in ((ws, "MetasSLA"), (alias_ws, "ClientAliases")):
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="002664")
    table = Table(displayName=table_name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet.freeze_panes = "A2"
    for column_cells in sheet.columns:
        width = max(len(str(cell.value or "")) for cell in column_cells) + 2
        sheet.column_dimensions[column_cells[0].column_letter].width = min(width, 24)

for row in ws.iter_rows(min_row=2, min_col=5, max_col=6):
    for cell in row:
        cell.number_format = "0.0%"

wb.save(OUTPUT)
print(OUTPUT)
